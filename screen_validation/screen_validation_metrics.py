#!/usr/bin/env python3
"""
screen_validation_metrics.py — Relevance-screen validation (NHSJS Rec #5)

Compares the LLM news-relevance screen's PASS/REJECT decisions against a blind
human labeling of the same articles, and reports precision, false-omission rate,
recall (re-weighted to the true population pass rate), F1, and Cohen's kappa,
with confidence intervals and a manuscript-ready sentence.

WHY RE-WEIGHTING: the labeling sample is STRATIFIED — it over-samples
screen-passed articles (70 passed + 110 rejected) so precision is estimated
stably. Because the two strata are sampled at different rates, RECALL and KAPPA
must be re-weighted to the population (pass rate 8.0% = 110/1382). PRECISION and
FALSE-OMISSION RATE condition on a single stratum, so they need no re-weighting.

INPUTS (both required):
  * The labeled sample workbook: screen_validation_sample.xlsx
      - sheet "Label" with columns: id, date, source, headline, snippet,
        "relevant? (1/0)"  (1 = human says relevant, 0 = not)
      - searched for in this folder, then the repo root.
  * screen_validation_key.csv  (the screen's blind decisions)
      - columns: id, screen_pass   (1 = screen PASSED, 0 = screen REJECTED)
      - This file is NOT generated here: it comes from the original screening
        run. See README_KEY.md in this folder. The script refuses to run on a
        missing or placeholder key rather than print meaningless numbers.

OPTIONAL:
  * _meta.json in this folder may override population counts:
        {"pop_pass": 110, "pop_total": 1382}

USAGE (after labeling every row and saving the xlsx):
    python screen_validation_metrics.py
"""
import csv
import json
import math
import os
import sys

# ---- population defaults (overridable via _meta.json) ----
POP_PASS = 110
POP_TOTAL = 1382  # POP_PASS_RATE = 110 / 1382 ~= 0.0796 (8.0%)

HERE = os.path.dirname(os.path.abspath(__file__))
REPO_ROOT = os.path.dirname(HERE)

XLSX_NAME = "screen_validation_sample.xlsx"
KEY_NAME = "screen_validation_key.csv"
LABEL_COL = "relevant? (1/0)"

N_BOOT = 20000
RNG_SEED = 20260727


def _find(name):
    for base in (HERE, REPO_ROOT):
        p = os.path.join(base, name)
        if os.path.exists(p):
            return p
    return None


def _load_meta():
    global POP_PASS, POP_TOTAL
    mp = _find("_meta.json")
    if mp:
        with open(mp) as f:
            m = json.load(f)
        POP_PASS = int(m.get("pop_pass", POP_PASS))
        POP_TOTAL = int(m.get("pop_total", POP_TOTAL))


def _load_labels(path):
    try:
        import openpyxl
    except ImportError:
        sys.exit("Need openpyxl:  pip install openpyxl --break-system-packages")
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Label"] if "Label" in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    try:
        i_id = header.index("id")
    except ValueError:
        sys.exit("Label sheet has no 'id' column.")
    # tolerant match for the label column
    i_lab = next((k for k, h in enumerate(header)
                  if h.lower().startswith("relevant")), None)
    if i_lab is None:
        sys.exit(f"Label sheet has no '{LABEL_COL}' column.")
    labels, unlabeled = {}, []
    for r in rows[1:]:
        if r[i_id] is None:
            continue
        aid = str(r[i_id]).strip()
        v = r[i_lab]
        if v is None or str(v).strip() == "":
            unlabeled.append(aid)
            continue
        labels[aid] = int(float(v))
    return labels, unlabeled


def _load_key(path):
    key = {}
    with open(path, newline="") as f:
        rdr = csv.DictReader(f)
        cols = {c.lower().strip(): c for c in (rdr.fieldnames or [])}
        if "id" not in cols:
            sys.exit(f"{KEY_NAME} needs an 'id' column.")
        pass_col = next((cols[c] for c in cols
                         if c in ("screen_pass", "llm_decision", "pass",
                                  "passed", "decision")), None)
        if pass_col is None:
            sys.exit(f"{KEY_NAME} needs a 'screen_pass' (1/0) column.")
        for row in rdr:
            aid = str(row[cols["id"]]).strip()
            raw = str(row[pass_col]).strip().lower()
            key[aid] = 1 if raw in ("1", "pass", "passed", "true", "yes") else 0
    return key


def wilson(k, n, z=1.96):
    if n == 0:
        return (float("nan"), float("nan"))
    p = k / n
    d = 1 + z * z / n
    c = p + z * z / (2 * n)
    h = z * math.sqrt((p * (1 - p) + z * z / (4 * n)) / n)
    return ((c - h) / d, (c + h) / d)


def pct(x):
    return "nan" if x != x else f"{100*x:.1f}%"


def main():
    _load_meta()
    xlsx = _find(XLSX_NAME)
    if not xlsx:
        sys.exit(f"Could not find {XLSX_NAME} in {HERE} or {REPO_ROOT}.")
    keyp = _find(KEY_NAME)
    if not keyp:
        sys.exit(
            f"Could not find {KEY_NAME}.\n"
            "This is the screen's actual PASS/REJECT decisions and cannot be "
            "reconstructed here. Restore it from the original screening run (or "
            "the repo-add archive) into this folder. See README_KEY.md.")

    labels, unlabeled = _load_labels(xlsx)
    key = _load_key(keyp)

    if unlabeled:
        print(f"WARNING: {len(unlabeled)} row(s) unlabeled and excluded: "
              f"{', '.join(unlabeled[:10])}{'...' if len(unlabeled) > 10 else ''}\n")

    ids = [i for i in key if i in labels]
    missing = [i for i in key if i not in labels]
    if missing and not unlabeled:
        print(f"NOTE: {len(missing)} key id(s) not found among labels.\n")
    if not ids:
        sys.exit("No overlap between key ids and labeled ids — check id formats.")

    # confusion within strata (screen decision x human label)
    # passed stratum
    a = sum(1 for i in ids if key[i] == 1 and labels[i] == 1)  # TP
    b = sum(1 for i in ids if key[i] == 1 and labels[i] == 0)  # FP
    # rejected stratum
    c = sum(1 for i in ids if key[i] == 0 and labels[i] == 1)  # FN
    d = sum(1 for i in ids if key[i] == 0 and labels[i] == 0)  # TN
    n_pass, n_rej = a + b, c + d

    precision = a / n_pass if n_pass else float("nan")   # no reweighting (single stratum)
    for_rate = c / n_rej if n_rej else float("nan")      # false-omission rate (single stratum)

    # population re-weighting for recall / kappa
    pop_rej = POP_TOTAL - POP_PASS
    w_pass = (POP_PASS / n_pass) if n_pass else 0.0
    w_rej = (pop_rej / n_rej) if n_rej else 0.0
    TP, FP = a * w_pass, b * w_pass
    FN, TN = c * w_rej, d * w_rej
    recall = TP / (TP + FN) if (TP + FN) else float("nan")
    f1 = (2 * precision * recall / (precision + recall)
          if precision + recall else float("nan"))

    # Cohen's kappa on re-weighted (population-representative) cells
    N = TP + FP + FN + TN
    po = (TP + TN) / N
    p_screen_pass = (TP + FP) / N
    p_human_rel = (TP + FN) / N
    pe = p_screen_pass * p_human_rel + (1 - p_screen_pass) * (1 - p_human_rel)
    kappa = (po - pe) / (1 - pe) if (1 - pe) else float("nan")

    # CIs: Wilson for the single-stratum rates; stratified bootstrap for recall & kappa
    prec_ci = wilson(a, n_pass)
    for_ci = wilson(c, n_rej)

    import random
    rnd = random.Random(RNG_SEED)
    pass_ids = [i for i in ids if key[i] == 1]
    rej_ids = [i for i in ids if key[i] == 0]
    rec_bs, kap_bs = [], []
    for _ in range(N_BOOT):
        ps = [rnd.choice(pass_ids) for _ in pass_ids]
        rs = [rnd.choice(rej_ids) for _ in rej_ids]
        aa = sum(1 for i in ps if labels[i] == 1); bb = len(ps) - aa
        cc = sum(1 for i in rs if labels[i] == 1); dd = len(rs) - cc
        TPb, FPb = aa * w_pass, bb * w_pass
        FNb, TNb = cc * w_rej, dd * w_rej
        if TPb + FNb:
            rec_bs.append(TPb / (TPb + FNb))
        Nb = TPb + FPb + FNb + TNb
        pob = (TPb + TNb) / Nb
        psp = (TPb + FPb) / Nb; phr = (TPb + FNb) / Nb
        peb = psp * phr + (1 - psp) * (1 - phr)
        if 1 - peb:
            kap_bs.append((pob - peb) / (1 - peb))
    rec_bs.sort(); kap_bs.sort()

    def ci(v):
        lo = v[int(0.025 * len(v))]; hi = v[int(0.975 * len(v)) - 1]
        return lo, hi
    rec_ci = ci(rec_bs) if rec_bs else (float("nan"), float("nan"))
    kap_ci = ci(kap_bs) if kap_bs else (float("nan"), float("nan"))

    print("=" * 68)
    print("RELEVANCE-SCREEN VALIDATION (NHSJS Rec #5)")
    print("=" * 68)
    print(f"Population pass rate assumed: {POP_PASS}/{POP_TOTAL} = "
          f"{100*POP_PASS/POP_TOTAL:.1f}%")
    print(f"Labeled sample: {len(ids)} articles "
          f"({n_pass} screen-passed, {n_rej} screen-rejected)\n")
    print("Confusion (rows = screen decision, cols = human label):")
    print(f"                 human=relevant   human=not")
    print(f"  screen=PASS         {a:>5}          {b:>5}")
    print(f"  screen=REJECT       {c:>5}          {d:>5}\n")
    print(f"Precision            {pct(precision)}  (95% CI {pct(prec_ci[0])}-{pct(prec_ci[1])})")
    print(f"False-omission rate  {pct(for_rate)}  (95% CI {pct(for_ci[0])}-{pct(for_ci[1])})")
    print(f"Recall (reweighted)  {pct(recall)}  (95% CI {pct(rec_ci[0])}-{pct(rec_ci[1])})")
    print(f"F1                   {pct(f1)}")
    print(f"Cohen's kappa        {kappa:.2f}  (95% CI {kap_ci[0]:.2f}-{kap_ci[1]:.2f})\n")

    # Landis-Koch descriptor for kappa
    def kappa_word(k):
        if k != k:
            return "undefined"
        if k < 0.0:
            return "poor (worse than chance)"
        if k < 0.20:
            return "slight"
        if k < 0.40:
            return "fair"
        if k < 0.60:
            return "moderate"
        if k < 0.80:
            return "substantial"
        return "almost perfect"

    print("Manuscript-ready sentence:")
    print("-" * 68)
    print(
        f"To validate the news-relevance screen, we blind-labeled a stratified "
        f"sample of {len(ids)} articles ({n_pass} screen-passed, {n_rej} "
        f"screen-rejected) drawn from the {POP_TOTAL}-article pool "
        f"(population pass rate {100*POP_PASS/POP_TOTAL:.1f}%). Against these "
        f"human labels the screen achieved a precision of {pct(precision)} "
        f"(95% CI {pct(prec_ci[0])}-{pct(prec_ci[1])}) and, re-weighting to the "
        f"population pass rate, an estimated recall of {pct(recall)} "
        f"(95% CI {pct(rec_ci[0])}-{pct(rec_ci[1])}); its false-omission rate "
        f"was {pct(for_rate)} and agreement with the human labels was "
        f"{kappa_word(kappa)} (Cohen's kappa = {kappa:.2f}).")
    print("-" * 68)


if __name__ == "__main__":
    main()
