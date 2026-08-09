#!/usr/bin/env python3
"""
portfolio_sim_permutation_runner.py — full permutation runner + GUI for the
cash-aware portfolio simulator (portfolio_sim_runner.py).

Feature parity with permutation_runner.py, adapted to the simulator's parameters,
PLUS per-permutation log viewing and per-worker live pop-out windows.

Parameter space (each dimension has include/exclude checkboxes at Start)
------------------------------------------------------------------------
  Sector          : technology | financials | healthcare
  Model           : gpt-4o | gpt-4o-mini | gpt-5.1
  Period length   : 12 mo | 24 mo
  Start month     : Jun 2024 -> Jun 2025 (13 monthly windows)
  Rebalance freq  : 2 | 4 | 6 | 12 segments per window
  Filter config   : X->Y  (X Stage-One candidates, Y = holdings cap)

Global run settings
-------------------
  OpenAI API key (saved to config), NYT key, results folder, budget, share mode
  (fractional / whole), risk-free rate, threads, runs-per-permutation, dry-run.

Controls
--------
  Start / Pause / Stop / Retry-errors, live progress + stats, sortable colored
  status table, per-run JSON output, "Save Excel now", session resume.

Logs (the point of the GUI)
---------------------------
  Each permutation's full simulator output is captured to its own buffer via a
  ThreadRoutedStream (thread-id routing => no interleaving across workers). Click
  a row to view its log; each worker also gets a live pop-out window.

Run modes
---------
  dry-run (default): MockBackend, NO OpenAI calls. --live uses RealBackend.

  GUI:       python portfolio_sim_permutation_runner.py
  Headless:  python portfolio_sim_permutation_runner.py --headless --out grid.json
  Excel:     python portfolio_sim_permutation_runner.py --save-excel-from-json
"""

from __future__ import annotations

import argparse
import itertools
import json
import os
import queue
import re
import sys
import threading
import time
import traceback
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass
from datetime import datetime
from typing import Any, Callable, Dict, List, Optional, Tuple

from dateutil.relativedelta import relativedelta

# --- load API keys from .env (kept out of git); safe no-op if .env is absent ---
import os as _os, pathlib as _pl
for _d in [_pl.Path(__file__).resolve().parent, *_pl.Path(__file__).resolve().parents]:
    _env = _d / ".env"
    if _env.exists():
        for _line in _env.read_text(encoding="utf-8").splitlines():
            _line = _line.strip()
            if _line and not _line.startswith("#") and "=" in _line:
                _k, _v = _line.split("=", 1)
                _os.environ.setdefault(_k.strip(), _v.strip().strip(chr(34)).strip(chr(39)))
        break
# --- end .env loader ---

ROOT = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, "SCCUR_TESTS"))

import sector_config  # noqa: E402

# ── Paths / config ─────────────────────────────────────────────────────────
RESULTS_DIR   = os.path.join(ROOT, "results", "portfolio_sim")
CONFIG_PATH   = os.path.join(ROOT, ".portfolio_sim_perm_config.json")
EXCEL_PATH    = os.path.join(RESULTS_DIR, "portfolio_sim_permutations.xlsx")
os.makedirs(RESULTS_DIR, exist_ok=True)

# ── Parameter space ────────────────────────────────────────────────────────
SECTORS        = ["technology", "financials", "healthcare"]
MODELS         = ["gpt-4o", "gpt-4o-mini", "gpt-5.1"]
PERIOD_MONTHS  = [12, 24]
FREQS          = [2, 4, 6, 12]                         # rebalance segments per window
FILTER_CONFIGS = [(20, 10), (20, 5), (10, 5), (10, 3), (5, 3), (30, 10), (30, 5)]

# Dimensions checked ON by default (others available but unchecked)
DEFAULT_SECTORS = {"technology"}
DEFAULT_MODELS  = {"gpt-5.1"}
DEFAULT_PERIODS = {12}
DEFAULT_FREQS   = set(FREQS)
DEFAULT_FILTERS = {(20, 10), (20, 5), (10, 5)}

_base       = datetime(2024, 6, 1)
START_DATES = [_base + relativedelta(months=i) for i in range(13)]   # Jun 2024 - Jun 2025

TRAINING_CUTOFFS = {
    "gpt-4o":      datetime(2023, 10, 1),
    "gpt-4o-mini": datetime(2023, 10, 1),
    "gpt-5":       datetime(2024, 10, 1),
    "gpt-5.1":     datetime(2024,  9, 1),
}
TODAY = datetime.now().replace(hour=0, minute=0, second=0, microsecond=0)

ST_PENDING = "Pending"
ST_SKIPPED = "Skipped"
ST_RUNNING = "Running"
ST_DONE    = "Done"
ST_ERROR   = "Error"
ROW_BG = {ST_PENDING: "", ST_SKIPPED: "#fff3cd", ST_RUNNING: "#cce5ff",
          ST_DONE: "#d4edda", ST_ERROR: "#f8d7da"}

BUDGET_DEFAULT = 10_000.0

_UNSET = object()   # sentinel: "no filter snapshot supplied" (distinct from None = invalid)


# ── Small helpers ──────────────────────────────────────────────────────────
def filter_label(X: int, Y: int, arrow: str = "->") -> str:
    return f"{X}{arrow}{Y}"


def parse_filter(cfg: str) -> Tuple[int, int]:
    for sep in ("->", ":", "/", "x", "X"):
        if sep in cfg:
            a, b = cfg.split(sep, 1)
            return int(a.strip()), int(b.strip())
    raise ValueError(f"bad filter config {cfg!r}; expected e.g. '20->10'")


def parse_filter_loose(cfg) -> Tuple[int, int]:
    """parse_filter, but also accepts '20-10' and [20, 10] — the forms people
    naturally write in a preset JSON."""
    if isinstance(cfg, (list, tuple)):
        return int(cfg[0]), int(cfg[1])
    cfg = str(cfg).strip()
    try:
        return parse_filter(cfg)
    except ValueError:
        if "-" in cfg:
            a, b = cfg.split("-", 1)
            return int(a.strip()), int(b.strip())
        raise


def _skip_reason(start: datetime, months: int, model: str) -> Optional[str]:
    end = start + relativedelta(months=months)
    cutoff = TRAINING_CUTOFFS.get(model)
    if cutoff and start <= cutoff:
        # Cancel entirely: the model's training cutoff is on/after the period start,
        # so it would have seen data from (or beyond) the decision date -> leakage.
        return f"Training cutoff ({model} {cutoff:%Y-%m}) >= start {start:%Y-%m}"
    if end > TODAY:
        return f"Period end {end:%Y-%m} > today ({TODAY:%Y-%m})"
    return None


def _safe_results_subdir(value: str) -> str:
    value = (value or "").strip().replace("\\", "/")
    parts = []
    for part in value.split("/"):
        clean = re.sub(r"[^A-Za-z0-9._-]+", "_", part.strip()).strip("._")
        if clean:
            parts.append(clean)
    return "/".join(parts)


def load_app_config() -> dict:
    try:
        with open(CONFIG_PATH, "r", encoding="utf-8") as f:
            data = json.load(f)
        return data if isinstance(data, dict) else {}
    except Exception:
        return {}


def save_app_config(cfg: dict) -> None:
    try:
        with open(CONFIG_PATH, "w", encoding="utf-8") as f:
            json.dump(cfg, f, indent=2)
    except Exception:
        pass


# ── Permutation model ──────────────────────────────────────────────────────
# ── Preset files ───────────────────────────────────────────────────────────
# A preset is a JSON file that preloads BOTH the run filters (which cells to
# run) and the option knobs (feedback, screened news, order mode, ...). It is
# applied at startup so the GUI opens already configured; you still press Start,
# so batch runs keep going through the UI as the repo rules require.
PRESET_OPTION_KEYS = (
    "threads", "runs_per_perm", "budget", "risk_free", "whole_shares", "dry_run",
    "results_subdir", "feedback", "txn_cost_bps", "screened_news", "min_articles",
    "order_mode", "min_articles_days", "rebalance_mode", "prompt_framing",
    "screened_corpus",
)

# Options that DEFINE the experiment. If a preset omits one of these it is forced
# back to the documented default rather than inheriting whatever the GUI/config
# happened to hold. Learned the hard way on 2026-07-31: running the hold-winners
# arm persisted rebalance_mode="hold" into .portfolio_sim_perm_config.json, and the
# feedback ON/OFF presets — which never mention rebalance_mode — silently inherited
# it. 108 runs came out in the wrong execution mode and had to be discarded.
EXPERIMENT_DEFINING_DEFAULTS = {
    "feedback": True,
    "screened_news": False,
    "order_mode": "dollars",
    "rebalance_mode": "reset",
    "prompt_framing": "reset",
    "screened_corpus": "",
    "txn_cost_bps": 0.0,
    "min_articles": 0,
    "min_articles_days": 90,
}


def load_preset(path: str) -> dict:
    """Read a preset JSON. Raises with a clear message if it is unusable."""
    with open(path, encoding="utf-8") as f:
        preset = json.load(f)
    if not isinstance(preset, dict):
        raise ValueError(f"{path}: preset must be a JSON object")
    unknown = set(preset) - {"name", "description", "options", "filters", "notes"}
    if unknown:
        print(f"  [preset] ignoring unknown top-level keys: {sorted(unknown)}")
    return preset


def apply_preset_to_cfg(preset: dict, cfg: dict) -> dict:
    """Fold a preset's `options` into the saved app-config dict."""
    opts = preset.get("options") or {}
    for k, v in opts.items():
        if k in PRESET_OPTION_KEYS:
            cfg[k] = v
        else:
            print(f"  [preset] ignoring unknown option {k!r}")
    return cfg


def build_permutations() -> List[dict]:
    rows, idx = [], 0
    for sector in SECTORS:
        for period in PERIOD_MONTHS:
            for start in START_DATES:
                for model in MODELS:
                    for freq in FREQS:
                        for (X, Y) in FILTER_CONFIGS:
                            idx += 1
                            end = start + relativedelta(months=period)
                            skip = _skip_reason(start, period, model)
                            rows.append({
                                "id": idx, "sector": sector, "period": period,
                                "start": start, "end": end, "model": model,
                                "freq": freq, "X": X, "Y": Y,
                                "skip": skip,
                                "status": ST_SKIPPED if skip else ST_PENDING,
                                "note": skip or "",
                                "result": None, "results": [], "elapsed_s": 0,
                                "log": [], "worker": None,
                            })
    return rows


_NEWS_DATES_CACHE: "Dict[tuple, object]" = {}


def _news_dates(screened: bool, corpus: str = ""):
    """Publication dates of the corpus a run would read (cached per process).

    `corpus` names an alternate screened corpus (e.g. tech_scored_ge7.csv) so the
    density gate measures the SAME file Stage One will read. Measuring the gate
    on one corpus while the run reads another would silently mis-gate cells.
    """
    import glob
    import pandas as pd
    key = (corpus or ("screened" if screened else "raw"),)
    if key in _NEWS_DATES_CACHE:
        return _NEWS_DATES_CACHE[key]
    if screened:
        paths = [os.path.join(ROOT, corpus or "tech_screened_corpus.csv")]
    else:
        paths = ([os.path.join(ROOT, "merged_news_data.csv")]
                 + sorted(glob.glob(os.path.join(ROOT, "tech_news_backfill_*.csv"))))
    frames = [pd.read_csv(p, usecols=["pub_date"]) for p in paths if os.path.exists(p)]
    if not frames:
        _NEWS_DATES_CACHE[key] = None
        return None
    d = pd.concat(frames, ignore_index=True)
    dates = pd.to_datetime(d.pub_date, errors="coerce").dropna().sort_values()
    _NEWS_DATES_CACHE[key] = dates
    return dates


def min_recent_articles(p: dict, screened: bool, days: int = 90,
                        corpus: str = "") -> "Optional[int]":
    """Fewest articles any of this permutation's decisions sees in its last `days`.

    The 2-year lookback never discriminates (every window clears ~470 articles);
    recency does, because a news-driven strategy reacts to what is new. Returns
    None when no corpus is available to measure.
    """
    import pandas as pd
    import portfolio_sim_runner as psr
    dates = _news_dates(screened, corpus)
    if dates is None or len(dates) == 0:
        return None
    segs = psr.build_segments(p["start"].strftime("%Y-%m-%d"),
                              p["end"].strftime("%Y-%m-%d"), n_segments=p["freq"])
    worst = None
    for s in segs:
        end = pd.Timestamp(s["start"])
        n = int(((dates > end - pd.DateOffset(days=days)) & (dates <= end)).sum())
        worst = n if worst is None else min(worst, n)
    return worst


def perm_label(p: dict) -> str:
    return (f"#{p['id']:04d} {p['sector'][:4]} {p['period']}mo "
            f"{p['start']:%Y-%m}->{p['end']:%Y-%m} {p['model']} "
            f"{p['freq']}x {filter_label(p['X'], p['Y'])}")


def perm_passes_filter(p: dict, f: dict) -> bool:
    if not (f["id_from"] <= p["id"] <= f["id_to"]):
        return False
    if p["sector"] not in f["sectors"]:
        return False
    if p["model"] not in f["models"]:
        return False
    if p["period"] not in f["periods"]:
        return False
    if p["freq"] not in f["freqs"]:
        return False
    if (p["X"], p["Y"]) not in f["filters"]:
        return False
    if not (f["date_from"] <= p["start"] <= f["date_to"]):
        return False
    return True


# ── Session resume (match saved JSON back to permutations by metadata) ──────
def _iter_result_json(subdir: str = "") -> List[str]:
    base = os.path.join(RESULTS_DIR, _safe_results_subdir(subdir)) if subdir else RESULTS_DIR
    out = []
    for root, _dirs, files in os.walk(base):
        for fn in files:
            if fn.startswith("psperm_") and fn.endswith(".json"):
                out.append(os.path.join(root, fn))
    return sorted(out)


def _match_perm(md: dict, by_id: dict) -> Optional[dict]:
    p = by_id.get(md.get("perm_id"))
    if p is None:
        return None
    # verify the core params match so id-shifts don't corrupt the restore
    try:
        if (p["sector"] == md.get("sector") and p["model"] == md.get("model")
                and p["period"] == md.get("period_months") and p["freq"] == md.get("freq")
                and p["X"] == md.get("filter_X") and p["Y"] == md.get("filter_Y")
                and p["start"].strftime("%Y-%m-%d") == md.get("analysis_start_date")):
            return p
    except Exception:
        return None
    return None


def result_overspend_errored(res: dict) -> bool:
    """True if a saved run overspent OVERSPEND_STREAK_LIMIT times in a row within any
    period. Runs now abort as ERROR when this happens; older saved runs instead fell
    back to a safe execution and were recorded as complete, so we re-flag them here."""
    import portfolio_sim_runner as psr
    for period in (res.get("periods") or []):
        streak = 0
        for reason in (period.get("rejected_reasons") or []):
            if psr._is_overspend_reason(reason):
                streak += 1
                if streak >= psr.OVERSPEND_STREAK_LIMIT:
                    return True
            else:
                streak = 0
    return False


def load_json_results_into_perms(perms: List[dict], subdir: str = "") -> Tuple[int, int]:
    """Restore prior runs so Start tops a cell up to runs-per-perm instead of
    re-running it.

    `subdir` MUST match the folder the current batch writes to. Scanning the
    results root while writing to an arm folder counts unrelated runs — e.g. the
    90 legacy runs would make a fresh feedback arm think each cell was already
    one-third complete, and it would schedule 2 runs instead of 3.
    """
    by_id = {p["id"]: p for p in perms}
    files_read = runs_loaded = 0
    for fp in _iter_result_json(subdir):
        files_read += 1
        try:
            with open(fp, "r", encoding="utf-8") as f:
                res = json.load(f)
        except Exception:
            continue
        if not isinstance(res, dict) or res.get("error") or "risk_metrics" not in res:
            continue
        # Mock runs must never count as progress. They are tagged two ways
        # (metadata flag since 2026-07-30, `dry_` filename prefix); honour both,
        # or a free dry preflight into an arm folder silently cancels the live
        # arm by making every cell look already-complete.
        if (res.get("metadata", {}) or {}).get("dry_run") \
                or "_dry_fb" in os.path.basename(fp):
            continue
        p = _match_perm(res.get("metadata", {}), by_id)
        if p is None or p["status"] == ST_SKIPPED:
            continue
        p["results"].append(res)
        p["result_path"] = fp
        runs_loaded += 1
    for p in perms:
        if p["results"]:
            p["result"] = p["results"][-1]
            p["elapsed_s"] = p["result"].get("metadata", {}).get("elapsed_s", 0)
            if result_overspend_errored(p["result"]):
                p["status"] = ST_ERROR
                p["note"] = "Overspent 3x in a row (marked error)"
            else:
                p["status"] = ST_DONE
    return files_read, runs_loaded


# ── Excel export ───────────────────────────────────────────────────────────
EXCEL_COLS = ["Perm ID", "Run #", "Sector", "Model", "Period (mo)", "Start", "End",
              "Rebalances", "Filter", "Budget", "Final Value", "Total ROI %",
              "S&P 500 %", "vs S&P %",
              "No-Reinvest ROI %", "vs No-Reinvest pp",
              "Ann. Return %", "Volatility %", "Sharpe", "Sortino", "Max DD %",
              "Calmar", "Beta", "Alpha %", "Rebalance events", "Reprompts"]


def _excel_rows_for_perm(p: dict) -> List[list]:
    rows = []
    for i, res in enumerate(p["results"], 1):
        rmx = res.get("risk_metrics") or {}
        def g(k, scale=1.0):
            v = rmx.get(k)
            return round(v * scale, 4) if isinstance(v, (int, float)) else None
        reprompts = sum(len(pd.get("rejected_reasons", [])) for pd in res.get("periods", []))
        _bench = res.get("benchmark_risk_metrics") or {}
        _spy = _bench.get("total_return")
        _spy_pct = round(_spy * 100, 4) if isinstance(_spy, (int, float)) else None
        _roi = res.get("total_roi_pct")
        _vs = (round(_roi - _spy_pct, 4)
               if isinstance(_roi, (int, float)) and _spy_pct is not None else None)
        _nri = res.get("no_reinvest_baseline") or {}
        _nri_roi = _nri.get("total_roi_pct")
        rows.append([
            p["id"], i, p["sector"], p["model"], p["period"],
            p["start"].strftime("%Y-%m-%d"), p["end"].strftime("%Y-%m-%d"),
            p["freq"] - 1, filter_label(p["X"], p["Y"]), res.get("budget"),
            res.get("final_value"), res.get("total_roi_pct"),
            _spy_pct, _vs,
            _nri_roi, res.get("vs_no_reinvest_pp"),
            g("annualized_return", 100), g("annualized_volatility", 100),
            g("sharpe"), g("sortino"), g("max_drawdown", 100), g("calmar"),
            g("beta"), g("alpha_annual", 100),
            res.get("num_rebalances"), reprompts,
        ])
    return rows


def save_excel(perms: List[dict], path: str) -> Optional[str]:
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        return "openpyxl not installed (pip install openpyxl)"
    global pd
    try:
        import pandas as pd  # noqa: F401  (used in _excel_rows_for_perm)
    except ImportError:
        pd = None
    wb = openpyxl.Workbook()
    by_sector: Dict[str, List[list]] = {}
    for p in perms:
        if p["results"]:
            by_sector.setdefault(p["sector"], []).extend(_excel_rows_for_perm(p))
    wb.remove(wb.active)
    hdr_fill = PatternFill("solid", fgColor="305496")
    for sector, rows in (by_sector or {"(none)": []}).items():
        ws = wb.create_sheet(sector[:28])
        ws.append(EXCEL_COLS)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = hdr_fill
        for r in sorted(rows, key=lambda r: (r[0], r[1])):
            ws.append(r)
        ws.freeze_panes = "A2"
        for col in ws.columns:
            w = max((len(str(c.value)) for c in col if c.value is not None), default=8)
            ws.column_dimensions[col[0].column_letter].width = min(max(w + 2, 8), 26)
    try:
        os.makedirs(os.path.dirname(path), exist_ok=True)
        wb.save(path)
    except Exception as e:
        return f"Excel save failed: {e}"
    return None


pd = None  # populated lazily by save_excel


# =========================================================================== #
#  Thread-routed stdout: keeps each permutation's log separate under threads   #
# =========================================================================== #
class ThreadRoutedStream:
    def __init__(self, real):
        self._real = real
        self._sinks: Dict[int, Callable[[str], None]] = {}
        self._lock = threading.Lock()

    def register(self, ident, sink):
        with self._lock:
            self._sinks[ident] = sink

    def unregister(self, ident):
        with self._lock:
            self._sinks.pop(ident, None)

    def write(self, s):
        sink = self._sinks.get(threading.get_ident())
        if sink is not None:
            sink(s)
        else:
            self._real.write(s)
        return len(s)

    def flush(self):
        try:
            self._real.flush()
        except Exception:
            pass


# =========================================================================== #
#  The single simulation seam (stubbable in tests; lazily imports simulator)   #
# =========================================================================== #
class _Stopped(Exception):
    """Raised inside a worker to abort an in-flight simulation the moment Stop is pressed."""


def _install_stop_checks(backend, stop_check):
    """Wrap the backend's data/decision calls so each one aborts if Stop was pressed.
    run_simulation calls these at every period, so an in-flight run halts at its next
    step (bounding wasted work to at most one already-issued API call)."""
    if not stop_check:
        return backend
    for name in ("get_candidates", "get_prices", "get_daily_prices", "decide"):
        orig = getattr(backend, name, None)
        if orig is None:
            continue

        def wrapped(*a, __orig=orig, **k):
            if stop_check():
                raise _Stopped()
            return __orig(*a, **k)
        setattr(backend, name, wrapped)
    return backend


def simulate_permutation(p: dict, *, api_key: str, dry_run: bool, budget: float,
                         whole_shares: bool, rf: float,
                         stop_check: "Optional[Callable[[], bool]]" = None,
                         feedback: "Optional[bool]" = None,
                         txn_cost_bps: "Optional[float]" = None,
                         screened_news: "Optional[bool]" = None,
                         order_mode: "Optional[str]" = None,
                         rebalance_mode: "Optional[str]" = None,
                         prompt_framing: "Optional[str]" = None,
                         screened_corpus: "Optional[str]" = None) -> dict:
    """Run one permutation through portfolio_sim_runner and return its result dict.

    feedback / txn_cost_bps default to the env vars PSIM_FEEDBACK (on|off) and
    PSIM_TXN_BPS so you can run the feedback-off control or a costed batch without
    touching the GUI: set PSIM_FEEDBACK=off (and a separate results folder to avoid
    filename collisions) before launching.
    """
    import hashlib
    import portfolio_sim_runner as psr
    if feedback is None:
        feedback = os.environ.get("PSIM_FEEDBACK", "on").strip().lower() != "off"
    if txn_cost_bps is None:
        txn_cost_bps = float(os.environ.get("PSIM_TXN_BPS", "0") or 0)
    if screened_news is None:
        screened_news = os.environ.get("PSIM_SCREENED_NEWS", "off").strip().lower() == "on"
    if order_mode is None:
        order_mode = os.environ.get("PSIM_ORDER_MODE", "dollars").strip().lower()
    if rebalance_mode is None:
        rebalance_mode = os.environ.get("PSIM_REBALANCE_MODE", "reset").strip().lower()

    X, Y = p["X"], p["Y"]
    print(f"{perm_label(p)}")
    print(f"  sector={p['sector']} model={p['model']} window={p['start']:%Y-%m-%d}->"
          f"{p['end']:%Y-%m-%d} segments={p['freq']} filter={X}->{Y} "
          f"budget=${budget:,.0f} {'whole' if whole_shares else 'fractional'} "
          f"{'DRY-RUN' if dry_run else 'LIVE'}")

    os.environ["SECTOR"] = p["sector"]
    segments = psr.build_segments(p["start"].strftime("%Y-%m-%d"),
                                  p["end"].strftime("%Y-%m-%d"), n_segments=p["freq"])
    print(f"  {len(segments)} holding segments, {len(segments)-1} rebalances")

    if dry_run:
        class _MockX(psr.MockBackend):
            def get_candidates(self, sector, as_of, max_holdings):
                u = sector_config.universe()
                h = int(hashlib.md5(f"cand{as_of}".encode()).hexdigest(), 16)
                k = min(len(u), X)
                return [u[(h + i * 5) % len(u)] for i in range(k)]
        backend = _MockX(order_mode=order_mode)
    else:
        if not api_key:
            raise RuntimeError("No OpenAI API key set (enter one in the GUI or set OPENAI_API_KEY).")
        os.environ["OPENAI_API_KEY"] = api_key
        class _RealX(psr.RealBackend):
            def get_candidates(self, sector, as_of, max_holdings):
                from dateutil.relativedelta import relativedelta as _rd
                from investment_strategy_generator import (_oai_create_with_retry,
                                                           extract_candidate_tickers)
                ns = psr._s(psr._d(as_of) - _rd(years=self.news_lookback_years))
                if self.screened_news:
                    # Paper B: LLM-relevance-screened corpus, and skip the sector
                    # substring filter (screened articles already passed a real test).
                    ndf = self._screened_window(ns, as_of)
                    print(f"   [screened news] {len(ndf)} articles in {ns}..{as_of}")
                else:
                    ndf = self.gen._load_news_data(sector_config.news_file(psr.ROOT), ns, as_of, sector)
                if ndf is None or len(ndf) == 0:
                    return []
                articles = self.gen._process_news_data(ndf, "" if self.screened_news else sector)
                model = self.gen._validate_model_for_date_range(as_of)
                prompt = (f"You are screening the {sector} sector as of {as_of}. Using ONLY the "
                          f"news below (on/before {as_of}), list the {X} strongest candidates.\n\n"
                          f"Ticker rules (STRICT): only common stocks listed on a U.S. exchange "
                          f"(NYSE, NASDAQ, NYSE American); use the exact U.S. root ticker with NO "
                          f"exchange suffix (no '.L', '.MI', '.DE', '.SW', '.PA', etc.); EXCLUDE "
                          f"ETFs/funds (including UCITS/foreign-domiciled), foreign ADRs, OTC names, "
                          f"and indices. Skip anything not directly investable as a U.S.-listed "
                          f"common stock.\n\nReturn ONLY JSON: a plain list of ticker symbols and "
                          f"nothing else (no rationale, no other fields): "
                          f'{{"candidates":["AAPL","MSFT"]}}\n\nNEWS:\n{articles}')
                # Retry a few times: a malformed/truncated screening response is usually
                # transient. If it still fails, skip this period (hold cash) instead of
                # crashing the whole run.
                for attempt in range(3):
                    try:
                        resp = _oai_create_with_retry(self.gen.openai_client, model=model,
                                                      messages=[{"role": "user", "content": prompt}],
                                                      **self.gen._temp_kwargs(model), max_completion_tokens=1500)
                        data = self.gen._load_json_response(resp.choices[0].message.content,
                                                            "candidate screening", model, max_completion_tokens=1500)
                        return extract_candidate_tickers(data)[:X]
                    except Exception as e:      # noqa: BLE001
                        print(f"   candidate screening attempt {attempt + 1}/3 failed: {e}")
                print("   candidate screening failed after 3 attempts; skipping this period (holding cash)")
                return []
        backend = _RealX(model=p["model"], screened_news=screened_news,
                         order_mode=order_mode, screened_corpus=screened_corpus)

    _install_stop_checks(backend, stop_check)
    run = psr.run_simulation(backend, sector=p["sector"], budget=budget, segments=segments,
                             max_holdings=Y, whole_shares=whole_shares, rf_annual=rf,
                             feedback=feedback, txn_cost_bps=txn_cost_bps,
                             rebalance_mode=rebalance_mode,
                             prompt_framing=prompt_framing or "reset")
    print(f"  feedback={'ON' if feedback else 'OFF'}  txn_cost={txn_cost_bps:g}bps  "
          f"screened_news={'ON' if screened_news else 'OFF'}  order_mode={order_mode}  "
          f"rebalance_mode={rebalance_mode}  prompt_framing={prompt_framing or 'reset'}")

    for pd_ in run["periods"]:
        extra = f"  (re-prompted {len(pd_['rejected_reasons'])}x)" if pd_.get("rejected_reasons") else ""
        print(f"  period {pd_['period']:>2} {pd_['stage']:>9} {pd_['as_of']}  "
              f"cash ${pd_['cash_after']:,.0f}  holds {pd_['num_holdings']}{extra}")
    rmx = run.get("risk_metrics") or {}
    print(f"  RESULT ROI {run['total_roi_pct']:+.2f}%  final ${run['final_value']:,.2f}  "
          f"Sharpe {rmx.get('sharpe')}  MaxDD {(rmx.get('max_drawdown') or 0)*100:.2f}%")

    run.setdefault("metadata", {}).update({
        "perm_id": p["id"], "sector": p["sector"], "model": p["model"],
        "period_months": p["period"], "freq": p["freq"], "filter_X": X, "filter_Y": Y,
        "analysis_start_date": p["start"].strftime("%Y-%m-%d"),
        "analysis_end_date": p["end"].strftime("%Y-%m-%d"),
    })
    return run


# module-level seam so tests can substitute a stub
SIMULATE_FN: Callable[..., dict] = simulate_permutation


# =========================================================================== #
#  Runner thread (ThreadPoolExecutor; pause/stop; runs-per-perm; retry-errors) #
# =========================================================================== #
class RunnerThread(threading.Thread):
    def __init__(self, perms, q, *, api_key, dry_run, budget, whole_shares, rf,
                 stop_evt, pause_evt, n_threads, filters, target_runs,
                 retry_errors_only=False, results_subdir="",
                 feedback=True, txn_cost_bps=0.0, screened_news=False,
                 min_articles=0, min_articles_days=90, order_mode="dollars",
                 rebalance_mode="reset", prompt_framing="reset",
                 screened_corpus=""):
        super().__init__(daemon=True)
        self.perms = perms
        self.q = q
        self.api_key = api_key
        self.dry_run = dry_run
        self.budget = budget
        self.whole_shares = whole_shares
        self.rf = rf
        self.stop_evt = stop_evt
        self.pause_evt = pause_evt
        self.n_threads = n_threads
        self.filters = filters
        self.target_runs = target_runs
        self.retry_errors_only = retry_errors_only
        self.feedback = bool(feedback)
        self.txn_cost_bps = float(txn_cost_bps or 0.0)
        self.screened_news = bool(screened_news)
        self.order_mode = order_mode or "dollars"
        self.rebalance_mode = rebalance_mode or "reset"
        self.prompt_framing = prompt_framing or "reset"
        self.screened_corpus = screened_corpus or ""
        self.min_articles = int(min_articles or 0)
        self.min_articles_days = int(min_articles_days or 90)
        self.results_subdir = _safe_results_subdir(results_subdir)
        self.output_dir = (os.path.join(RESULTS_DIR, self.results_subdir)
                           if self.results_subdir else RESULTS_DIR)
        self._router: Optional[ThreadRoutedStream] = None
        self._saved_stdout = None
        self._worker_ids: Dict[int, int] = {}
        self._worker_lock = threading.Lock()

    def _post(self, kind, **kw):
        self.q.put({"kind": kind, **kw})

    def _worker_index(self) -> int:
        ident = threading.get_ident()
        with self._worker_lock:
            if ident not in self._worker_ids:
                self._worker_ids[ident] = len(self._worker_ids)
            return self._worker_ids[ident]

    def _run_one(self, task):
        p = task["perm"]
        if self.stop_evt.is_set():
            return                 # Stop pressed before this task started any work
        log = task["log"]          # this run's OWN buffer (isolates repeated/concurrent runs)
        p["log"] = log             # row displays the latest run's log
        widx = self._worker_index()
        p["worker"] = widx
        run_token = f"{p['id']}-{threading.get_ident()}-{time.time_ns()}"

        # Pre-flight news check: refuse cells whose thinnest decision cannot see
        # enough recent news to make a real choice. Done BEFORE any API call so a
        # starved cell costs nothing, and marked Skipped so it stays visible in
        # the table and the run ledger rather than silently vanishing.
        if self.min_articles > 0:
            worst = min_recent_articles(p, self.screened_news, self.min_articles_days,
                                    self.screened_corpus)
            if worst is not None and worst < self.min_articles:
                note = (f"Skipped: thinnest decision sees {worst} screened articles "
                        f"in {self.min_articles_days}d (min {self.min_articles})")
                self._post("status", id=p["id"], run_token=run_token,
                           status=ST_SKIPPED, note=note, worker=widx)
                self._post("log", text=f"[{datetime.now():%H:%M:%S}] - #{p['id']:04d} {note}")
                return

        self._post("status", id=p["id"], run_token=run_token, status=ST_RUNNING,
                   note="", worker=widx)
        self._post("log", text=f"[{datetime.now():%H:%M:%S}] > {perm_label(p)}  (worker {widx})")

        ident = threading.get_ident()

        def sink(text, _log=log, _pid=p["id"], _w=widx):
            _log.append(text)
            self._post("perm_log", id=_pid, worker=_w, text=text)

        self._router.register(ident, sink)
        t0 = time.time()
        try:
            result = SIMULATE_FN(p, api_key=self.api_key, dry_run=self.dry_run,
                                 budget=self.budget, whole_shares=self.whole_shares, rf=self.rf,
                                 stop_check=self.stop_evt.is_set,
                                 feedback=self.feedback, txn_cost_bps=self.txn_cost_bps,
                                 screened_news=self.screened_news,
                                 order_mode=self.order_mode,
                                 rebalance_mode=self.rebalance_mode,
                                 prompt_framing=self.prompt_framing,
                                 screened_corpus=self.screened_corpus)
            elapsed = round(time.time() - t0)
            # _run_one is authoritative for the identifying metadata so session
            # resume matches regardless of what the simulate function returned.
            result.setdefault("metadata", {}).update({
                "perm_id": p["id"], "sector": p["sector"], "model": p["model"],
                "period_months": p["period"], "freq": p["freq"],
                "filter_X": p["X"], "filter_Y": p["Y"],
                "analysis_start_date": p["start"].strftime("%Y-%m-%d"),
                "analysis_end_date": p["end"].strftime("%Y-%m-%d"),
                "elapsed_s": elapsed,
                "dry_run": self.dry_run,
            })
            # write per-run JSON
            ts = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
            # Encode the arm in the filename so ON/OFF (and costed) runs stay
            # distinguishable even if they ever land in the same folder.
            arm = "dry_" if self.dry_run else ""
            arm += "fbon" if self.feedback else "fboff"
            if self.txn_cost_bps:
                arm += f"_{self.txn_cost_bps:g}bps"
            if self.screened_news:
                arm += "_scr"
            if self.order_mode == "weights":
                arm += "_wt"
            if self.rebalance_mode == "hold":
                arm += "_hold"
            elif self.rebalance_mode == "hold_redeploy":
                arm += "_holdrd"
            if self.screened_corpus:
                m = re.search(r"ge(\d+)", self.screened_corpus)
                arm += "_ge%s" % m.group(1) if m else "_alt"
            if self.prompt_framing == "incumbent":
                arm += "_inc"
            elif self.prompt_framing == "incumbent_invested":
                arm += "_incinv"
            fname = (f"psperm_{p['id']:04d}_{p['sector'][:4]}_{p['model'].replace('.', '')}_"
                     f"{p['freq']}x_{p['X']}-{p['Y']}_{p['start']:%Y%m}_{ts}_{arm}_w{widx}.json")
            os.makedirs(self.output_dir, exist_ok=True)
            json_path = os.path.join(self.output_dir, fname)
            with open(json_path, "w", encoding="utf-8") as f:
                json.dump(result, f, indent=2, default=str)
            rel_fname = os.path.join(self.results_subdir, fname) if self.results_subdir else fname
            roi = result.get("total_roi_pct")
            rmx = result.get("risk_metrics") or {}
            vsn = result.get("vs_no_reinvest_pp")
            vsn_txt = f"  vsNOREINV {vsn:+.2f}pp" if isinstance(vsn, (int, float)) else ""
            note = (f"ROI {roi:+.2f}%  Sharpe {rmx.get('sharpe')}{vsn_txt}  "
                    f"{elapsed}s  |  {rel_fname}") if roi is not None else f"{elapsed}s | {rel_fname}"
            self._post("status", id=p["id"], run_token=run_token, status=ST_DONE,
                       note=note, result=result, elapsed_s=elapsed, worker=widx,
                       json_path=json_path)
            self._post("log", text=f"[{datetime.now():%H:%M:%S}] v #{p['id']:04d} {note}")
        except _Stopped:
            # Stop pressed mid-run: revert to Pending (resumable), no JSON, no error.
            self._post("status", id=p["id"], run_token=run_token, status=ST_PENDING,
                       note="Stopped", result=None, worker=widx)
            self._post("log", text=f"[{datetime.now():%H:%M:%S}] - #{p['id']:04d} stopped")
        except Exception as exc:
            elapsed = round(time.time() - t0)
            tb = traceback.format_exc()
            sink("\n[ERROR]\n" + tb)
            self._post("status", id=p["id"], run_token=run_token, status=ST_ERROR,
                       note=f"{elapsed}s | {str(exc)[:120]}", result=None,
                       elapsed_s=elapsed, worker=widx)
            self._post("log", text=f"[{datetime.now():%H:%M:%S}] x #{p['id']:04d} ERROR: {exc}")
            try:
                logs_dir = os.path.join(self.output_dir, "logs")
                os.makedirs(logs_dir, exist_ok=True)
                with open(os.path.join(logs_dir, f"psperm_{p['id']:04d}_error_"
                          f"{datetime.now():%Y%m%d_%H%M%S}.txt"), "w", encoding="utf-8") as lf:
                    lf.write(perm_label(p) + "\n\n" + tb)
            except Exception:
                pass
        finally:
            self._router.unregister(ident)

    def run(self):
        self._saved_stdout = sys.stdout
        self._router = ThreadRoutedStream(self._saved_stdout)
        sys.stdout = self._router
        try:
            work = []
            if self.retry_errors_only:
                for p in self.perms:
                    if p["status"] == ST_ERROR and perm_passes_filter(p, self.filters):
                        work.append({"perm": p, "log": []})
            else:
                for p in self.perms:
                    if p["status"] == ST_SKIPPED or not perm_passes_filter(p, self.filters):
                        continue
                    for _ in range(max(0, self.target_runs - len(p.get("results", [])))):
                        work.append({"perm": p, "log": []})

            scheduled = {}
            for task in work:
                pid = task["perm"]["id"]
                scheduled[pid] = scheduled.get(pid, 0) + 1
            self._post("plan", scheduled_by_pid=scheduled, total=len(work))

            sem = threading.Semaphore(self.n_threads)
            with ThreadPoolExecutor(max_workers=self.n_threads) as ex:
                futures = []
                for task in work:
                    if self.stop_evt.is_set():
                        break
                    while self.pause_evt.is_set() and not self.stop_evt.is_set():
                        time.sleep(0.3)
                    if self.stop_evt.is_set():
                        break
                    sem.acquire()
                    if self.stop_evt.is_set():
                        sem.release()
                        break

                    def _task(tk=task):
                        try:
                            self._run_one(tk)
                        finally:
                            sem.release()
                    futures.append(ex.submit(_task))
                for f in futures:
                    try:
                        f.result()
                    except Exception:
                        pass
        finally:
            sys.stdout = self._saved_stdout
            self._post("done")


# =========================================================================== #
#  Tkinter GUI                                                                #
# =========================================================================== #
class App:
    def __init__(self, preset: "Optional[dict]" = None,
                 autostart: bool = False, autostart_delay: int = 8):
        import tkinter as tk
        from tkinter import ttk, scrolledtext, messagebox, filedialog
        self.tk = tk; self.ttk = ttk; self.scrolledtext = scrolledtext; self.messagebox = messagebox
        self.filedialog = filedialog

        self.root = tk.Tk()
        self.root.title("Portfolio-Sim Permutation Runner")
        self.root.geometry("1360x840")
        self.root.minsize(1000, 640)

        self.perms = build_permutations()
        self.cfg = load_app_config()
        # Preset options must land in cfg BEFORE the tk vars read their defaults.
        self.preset = preset
        self.autostart = bool(autostart)
        self.autostart_delay = max(0, int(autostart_delay))
        if preset:
            self.cfg = apply_preset_to_cfg(preset, self.cfg)
        self._q: "queue.Queue[dict]" = queue.Queue()
        self._stop = threading.Event()
        self._pause = threading.Event()
        self._worker: Optional[RunnerThread] = None
        self._iid: Dict[int, str] = {}
        self._sort_rev: Dict[str, bool] = {}
        self._worker_windows: Dict[int, dict] = {}
        self._run_total = self._run_finished = 0
        self._batch_sched: Dict[int, int] = {}   # perm_id -> runs scheduled in current batch
        self._batch_done: Dict[int, int] = {}    # perm_id -> runs completed in current batch
        # ---- list-view filters: which permutations appear in the main table ----
        self.view_status_vars = {st: tk.BooleanVar(value=True)
                                 for st in (ST_PENDING, ST_SKIPPED, ST_RUNNING, ST_DONE, ST_ERROR)}
        self.view_match_filters = tk.BooleanVar(value=True)   # honor the Run Filters in the list

        maxid = max(p["id"] for p in self.perms)
        self.id_from = tk.IntVar(value=1); self.id_to = tk.IntVar(value=maxid)
        self.sector_vars = {s: tk.BooleanVar(value=s in DEFAULT_SECTORS) for s in SECTORS}
        self.model_vars  = {m: tk.BooleanVar(value=m in DEFAULT_MODELS) for m in MODELS}
        self.period_vars = {p: tk.BooleanVar(value=p in DEFAULT_PERIODS) for p in PERIOD_MONTHS}
        self.freq_vars   = {fq: tk.BooleanVar(value=fq in DEFAULT_FREQS) for fq in FREQS}
        self.filt_vars   = {(X, Y): tk.BooleanVar(value=(X, Y) in DEFAULT_FILTERS) for (X, Y) in FILTER_CONFIGS}
        dstr = [d.strftime("%Y-%m") for d in START_DATES]
        self.date_from = tk.StringVar(value=dstr[0]); self.date_to = tk.StringVar(value=dstr[-1])

        self.threads_var = tk.IntVar(value=int(self.cfg.get("threads", 1)))
        self.runs_var    = tk.IntVar(value=int(self.cfg.get("runs_per_perm", 1)))
        self.budget_var  = tk.DoubleVar(value=float(self.cfg.get("budget", BUDGET_DEFAULT)))
        self.rf_var      = tk.DoubleVar(value=float(self.cfg.get("risk_free", 0.0)))
        self.whole_var   = tk.BooleanVar(value=bool(self.cfg.get("whole_shares", False)))
        self.dry_var     = tk.BooleanVar(value=bool(self.cfg.get("dry_run", True)))
        self.key_var     = tk.StringVar(value=self.cfg.get("openai_api_key") or os.environ.get("OPENAI_API_KEY", ""))
        self.nyt_var     = tk.StringVar(value=self.cfg.get("nyt_api_key", ""))
        self.folder_var  = tk.StringVar(value=self.cfg.get("results_subdir", ""))
        # Feedback (self-reflection loop) and transaction cost are native controls:
        # the ON/OFF experiment and the cost sweep run from here, never a side loop.
        self.fb_var      = tk.BooleanVar(value=bool(self.cfg.get("feedback", True)))
        self.txn_var     = tk.DoubleVar(value=float(self.cfg.get("txn_cost_bps", 0.0)))
        # Paper B: read the LLM-relevance-screened corpus instead of the raw merge.
        self.screen_var  = tk.BooleanVar(value=bool(self.cfg.get("screened_news", False)))
        # Refuse cells whose thinnest decision has too little recent news to matter.
        self.minart_var  = tk.IntVar(value=int(self.cfg.get("min_articles", 0)))
        self.minartdays_var = tk.IntVar(value=int(self.cfg.get("min_articles_days", 90)))
        # dollars = model emits dollar orders against a hard budget (original).
        # weights = model emits target fractions; the engine does the arithmetic.
        self.ordmode_var = tk.StringVar(value=self.cfg.get("order_mode", "dollars"))
        # reset = liquidate the whole book each rebalance and re-buy to fresh target
        #         weights (original). Resets a compounding winner back to ~1/N.
        # hold  = trade only the delta: sell what the model dropped, buy what it
        #         added, and leave kept positions alone so winners keep compounding.
        #         Requires Orders=weights. See execute_hold_mode() in the engine.
        self.rebmode_var = tk.StringVar(value=self.cfg.get("rebalance_mode", "reset"))
        # reset     = the rebalance prompt tells the model its book "has already been
        #             SOLD" and asks it to rebuild from scratch (original wording).
        # incumbent = the prompt states the book it actually still owns, gives the gap
        #             to the next decision, and sets a switching hurdle. Targets the
        #             measured -1.6..-2.6pp swap edge (dropped names outrun added ones)
        #             and the cadence inversion (2x ~ -4..-6.5pp/decision, 12x ~ 0).
        #             With feedback OFF the book is shown as STATE ONLY, so the control
        #             arm still receives no performance signal.
        self.framing_var = tk.StringVar(value=self.cfg.get("prompt_framing", "reset"))
        # Alternate screened corpus (tech_scored_ge4..ge8.csv). Empty = the
        # default tech_screened_corpus.csv. Only meaningful with Screened news on.
        self.corpus_var = tk.StringVar(value=self.cfg.get("screened_corpus", ""))

        self._filter_widgets: list = []
        self._build_ui()
        if self.preset:
            self._apply_preset_options(self.preset)
            self._apply_preset_filters(self.preset)
        if self.autostart:
            self.root.after(300, self._begin_autostart)
        self._load_existing()
        self._fill_table()
        self._wire_view_traces()
        self._refresh_stats()
        self.root.after(200, self._drain)

    # ---- UI ------------------------------------------------------------- #
    def _build_ui(self):
        tk, ttk = self.tk, self.ttk
        bar = ttk.Frame(self.root, padding=(8, 6)); bar.pack(fill="x")
        self.b_start = ttk.Button(bar, text="▶ Start", command=self._on_start)
        self.b_pause = ttk.Button(bar, text="⏸ Pause", command=self._on_pause, state="disabled")
        self.b_stop  = ttk.Button(bar, text="⏹ Stop", command=self._on_stop, state="disabled")
        self.b_retry = ttk.Button(bar, text="Retry Errors", command=self._on_retry)
        self.b_retry_all = ttk.Button(bar, text="Retry All (replace)", command=self._on_retry_all)
        self.b_excel = ttk.Button(bar, text="Save Excel now", command=self._on_excel)
        self.b_batch = ttk.Button(bar, text="Queued runs", command=self._reopen_batch_window)
        self.b_done_only = ttk.Button(bar, text="Show completed only",
                                      command=self._view_completed_only)
        self.b_preset = ttk.Button(bar, text="Load preset…", command=self._on_load_preset)
        for b in (self.b_start, self.b_pause, self.b_stop, self.b_retry, self.b_retry_all,
                  self.b_excel, self.b_batch, self.b_done_only, self.b_preset):
            b.pack(side="left", padx=3)
        ttk.Separator(bar, orient="vertical").pack(side="left", fill="y", padx=8)
        ttk.Label(bar, text="Threads:").pack(side="left")
        ttk.Spinbox(bar, from_=1, to=16, width=4, textvariable=self.threads_var).pack(side="left", padx=(2, 8))
        ttk.Label(bar, text="Runs/perm:").pack(side="left")
        ttk.Spinbox(bar, from_=1, to=50, width=4, textvariable=self.runs_var).pack(side="left", padx=(2, 8))
        ttk.Checkbutton(bar, text="Dry-run (no API)", variable=self.dry_var).pack(side="left", padx=6)
        self.stats_var = tk.StringVar()
        ttk.Label(bar, textvariable=self.stats_var, font=("Courier", 10)).pack(side="left", padx=8)

        # Row: API keys, budget, share mode, folder
        opt = ttk.Frame(self.root, padding=(8, 0)); opt.pack(fill="x")
        ttk.Label(opt, text="OpenAI key:").pack(side="left")
        ttk.Entry(opt, textvariable=self.key_var, width=40, show="*").pack(side="left", padx=(2, 8))
        ttk.Label(opt, text="NYT key:").pack(side="left")
        ttk.Entry(opt, textvariable=self.nyt_var, width=18, show="*").pack(side="left", padx=(2, 8))
        ttk.Label(opt, text="Budget $:").pack(side="left")
        ttk.Entry(opt, textvariable=self.budget_var, width=9).pack(side="left", padx=(2, 8))
        ttk.Label(opt, text="Risk-free:").pack(side="left")
        ttk.Entry(opt, textvariable=self.rf_var, width=6).pack(side="left", padx=(2, 8))
        ttk.Checkbutton(opt, text="Whole shares", variable=self.whole_var).pack(side="left", padx=6)
        self.fb_check = ttk.Checkbutton(opt, text="Feedback", variable=self.fb_var)
        self.fb_check.pack(side="left", padx=6)
        ttk.Label(opt, text="Txn bps:").pack(side="left")
        self.txn_entry = ttk.Entry(opt, textvariable=self.txn_var, width=6)
        self.txn_entry.pack(side="left", padx=(2, 8))
        self.screen_check = ttk.Checkbutton(opt, text="Screened news", variable=self.screen_var)
        self.screen_check.pack(side="left", padx=6)
        ttk.Label(opt, text="Orders:").pack(side="left")
        self.ordmode_combo = ttk.Combobox(opt, textvariable=self.ordmode_var, width=8,
                                          state="readonly", values=("dollars", "weights"))
        self.ordmode_combo.pack(side="left", padx=(2, 8))
        ttk.Label(opt, text="Rebalance:").pack(side="left")
        self.rebmode_combo = ttk.Combobox(opt, textvariable=self.rebmode_var, width=14,
                                          state="readonly",
                                          values=("reset", "hold", "hold_redeploy"))
        self.rebmode_combo.pack(side="left", padx=(2, 8))
        ttk.Label(opt, text="Framing:").pack(side="left")
        self.framing_combo = ttk.Combobox(opt, textvariable=self.framing_var, width=18,
                                          state="readonly",
                                          values=("reset", "incumbent", "incumbent_invested"))
        self.framing_combo.pack(side="left", padx=(2, 8))
        ttk.Label(opt, text="Min articles:").pack(side="left")
        self.minart_entry = ttk.Entry(opt, textvariable=self.minart_var, width=5)
        self.minart_entry.pack(side="left", padx=(2, 1))
        ttk.Label(opt, text="per days:").pack(side="left")
        self.minartdays_entry = ttk.Entry(opt, textvariable=self.minartdays_var, width=5)
        self.minartdays_entry.pack(side="left", padx=(2, 8))
        ttk.Label(opt, text="Folder:").pack(side="left")
        self.folder_entry = ttk.Entry(opt, textvariable=self.folder_var, width=16)
        self.folder_entry.pack(side="left", padx=(2, 8))

        # Progress + legend
        pbf = ttk.Frame(self.root, padding=(8, 2)); pbf.pack(fill="x")
        self.run_prog_var = tk.StringVar(value="Runs: — queued / — completed")
        ttk.Label(pbf, textvariable=self.run_prog_var, font=("Segoe UI", 9, "bold")).pack(anchor="w")
        self.pb_var = tk.DoubleVar(value=0.0)
        ttk.Progressbar(pbf, variable=self.pb_var, maximum=100).pack(fill="x", expand=True)
        self.cur_var = tk.StringVar(value="Not started")
        ttk.Label(pbf, textvariable=self.cur_var, font=("Courier", 9), foreground="#444").pack(anchor="w")
        leg = ttk.Frame(self.root, padding=(8, 0)); leg.pack(fill="x")
        for lab, col in [("Pending", "#e0e0e0"), ("Skipped", "#fff3cd"), ("Running", "#cce5ff"),
                         ("Done", "#d4edda"), ("Error", "#f8d7da")]:
            tk.Label(leg, text=f"  {lab}  ", background=col, relief="solid", bd=1,
                     font=("Courier", 8)).pack(side="left", padx=4, pady=2)

        # Filters
        fb = ttk.LabelFrame(self.root, text="Run Filters (applied at Start)", padding=(8, 4)); fb.pack(fill="x", padx=6)
        r1 = ttk.Frame(fb); r1.pack(fill="x", anchor="w")
        ttk.Label(r1, text="ID:").pack(side="left")
        w = ttk.Spinbox(r1, from_=1, to=99999, width=6, textvariable=self.id_from); w.pack(side="left", padx=2); self._filter_widgets.append(w)
        ttk.Label(r1, text="to").pack(side="left")
        w = ttk.Spinbox(r1, from_=1, to=99999, width=6, textvariable=self.id_to); w.pack(side="left", padx=2); self._filter_widgets.append(w)
        ttk.Separator(r1, orient="vertical").pack(side="left", fill="y", padx=8)
        ttk.Label(r1, text="Sectors:").pack(side="left")
        for s in SECTORS:
            w = ttk.Checkbutton(r1, text=s, variable=self.sector_vars[s]); w.pack(side="left", padx=2); self._filter_widgets.append(w)
        ttk.Separator(r1, orient="vertical").pack(side="left", fill="y", padx=8)
        ttk.Label(r1, text="Models:").pack(side="left")
        for m in MODELS:
            w = ttk.Checkbutton(r1, text=m, variable=self.model_vars[m]); w.pack(side="left", padx=2); self._filter_widgets.append(w)

        r2 = ttk.Frame(fb); r2.pack(fill="x", anchor="w", pady=(3, 0))
        ttk.Label(r2, text="Periods:").pack(side="left")
        for pr in PERIOD_MONTHS:
            w = ttk.Checkbutton(r2, text=f"{pr}mo", variable=self.period_vars[pr]); w.pack(side="left", padx=2); self._filter_widgets.append(w)
        ttk.Separator(r2, orient="vertical").pack(side="left", fill="y", padx=8)
        ttk.Label(r2, text="Rebalance freqs:").pack(side="left")
        for fq in FREQS:
            w = ttk.Checkbutton(r2, text=f"{fq}x", variable=self.freq_vars[fq]); w.pack(side="left", padx=2); self._filter_widgets.append(w)
        ttk.Separator(r2, orient="vertical").pack(side="left", fill="y", padx=8)
        ttk.Label(r2, text="Filters:").pack(side="left")
        for (X, Y) in FILTER_CONFIGS:
            w = ttk.Checkbutton(r2, text=filter_label(X, Y), variable=self.filt_vars[(X, Y)]); w.pack(side="left", padx=2); self._filter_widgets.append(w)

        r3 = ttk.Frame(fb); r3.pack(fill="x", anchor="w", pady=(3, 0))
        dstr = [d.strftime("%Y-%m") for d in START_DATES]
        ttk.Label(r3, text="Start date:").pack(side="left")
        w = ttk.Combobox(r3, textvariable=self.date_from, values=dstr, state="readonly", width=9); w.pack(side="left", padx=2); self._filter_widgets.append(w)
        ttk.Label(r3, text="to").pack(side="left")
        w = ttk.Combobox(r3, textvariable=self.date_to, values=dstr, state="readonly", width=9); w.pack(side="left", padx=2); self._filter_widgets.append(w)

        # View filter: which permutations appear in the list below
        vf = ttk.LabelFrame(self.root, text="Show in list (view filter)", padding=(8, 4)); vf.pack(fill="x", padx=6)
        vr = ttk.Frame(vf); vr.pack(fill="x", anchor="w")
        ttk.Label(vr, text="Status:").pack(side="left")
        for st in (ST_PENDING, ST_SKIPPED, ST_RUNNING, ST_DONE, ST_ERROR):
            ttk.Checkbutton(vr, text=st, variable=self.view_status_vars[st],
                            command=self._fill_table).pack(side="left", padx=2)
        ttk.Separator(vr, orient="vertical").pack(side="left", fill="y", padx=8)
        ttk.Checkbutton(vr, text="Match Run Filters above", variable=self.view_match_filters,
                        command=self._fill_table).pack(side="left", padx=2)
        ttk.Button(vr, text="Show all", command=self._reset_view).pack(side="left", padx=8)

        # Table + log
        pw = ttk.PanedWindow(self.root, orient="vertical"); pw.pack(fill="both", expand=True, padx=6, pady=4)
        tblf = ttk.Frame(pw); pw.add(tblf, weight=4)
        self.COLS = ("#", "Sector", "Period", "Start", "End", "Model", "Freq", "Filter",
                     "Status", "ROI", "S&P 500", "vs S&P", "Sharpe", "MaxDD", "Note")
        widths = {"#": 46, "Sector": 74, "Period": 52, "Start": 70, "End": 70, "Model": 84,
                  "Freq": 42, "Filter": 66, "Status": 64, "ROI": 70, "S&P 500": 70, "vs S&P": 66,
                  "Sharpe": 56, "MaxDD": 64, "Note": 300}
        self.tree = ttk.Treeview(tblf, columns=self.COLS, show="headings", selectmode="browse")
        for c in self.COLS:
            self.tree.heading(c, text=c, command=lambda cc=c: self._sort(cc))
            self.tree.column(c, width=widths[c], anchor="w" if c == "Note" else "center", stretch=(c == "Note"))
        ys = ttk.Scrollbar(tblf, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=ys.set); ys.pack(side="right", fill="y")
        self.tree.pack(fill="both", expand=True)
        for st, bg in ROW_BG.items():
            if bg:
                self.tree.tag_configure(st, background=bg)
        self.tree.bind("<<TreeviewSelect>>", self._show_log)
        self.tree.bind("<Double-1>", self._open_perm_json)

        logf = ttk.LabelFrame(pw, text="Log for selected permutation", padding=4); pw.add(logf, weight=2)
        self.log_box = self.scrolledtext.ScrolledText(logf, height=10, state="disabled",
                                                      font=("Consolas", 9), wrap="word")
        self.log_box.pack(fill="both", expand=True)
        gl = ttk.LabelFrame(pw, text="Activity log", padding=4); pw.add(gl, weight=1)
        self.act_box = self.scrolledtext.ScrolledText(gl, height=5, state="disabled",
                                                      font=("Consolas", 9), wrap="word")
        self.act_box.pack(fill="both", expand=True)

    # ---- filters snapshot ---------------------------------------------- #
    def _snapshot_filters(self) -> dict:
        df = datetime.strptime(self.date_from.get(), "%Y-%m")
        dt = datetime.strptime(self.date_to.get(), "%Y-%m")
        if df > dt:
            df, dt = dt, df
        return {
            "id_from": self.id_from.get(), "id_to": self.id_to.get(),
            "sectors": {s for s, v in self.sector_vars.items() if v.get()},
            "models": {m for m, v in self.model_vars.items() if v.get()},
            "periods": {p for p, v in self.period_vars.items() if v.get()},
            "freqs": {fq for fq, v in self.freq_vars.items() if v.get()},
            "filters": {k for k, v in self.filt_vars.items() if v.get()},
            "date_from": df, "date_to": dt,
        }

    # ---- table --------------------------------------------------------- #
    def _row_vals(self, p):
        r = p.get("result") or {}
        rmx = (r.get("risk_metrics") or {}) if r else {}
        bench = (r.get("benchmark_risk_metrics") or {}) if r else {}
        roi_v = r.get("total_roi_pct")
        spy_v = bench.get("total_return")                     # SPY return over the same window
        roi = "" if roi_v is None else f"{roi_v:+.2f}%"
        spy = "" if spy_v is None else f"{spy_v * 100:+.2f}%"
        vs = "" if (roi_v is None or spy_v is None) else f"{roi_v - spy_v * 100:+.2f}%"
        shp = "" if rmx.get("sharpe") is None else f"{rmx['sharpe']:.2f}"
        mdd = "" if rmx.get("max_drawdown") is None else f"{rmx['max_drawdown']*100:+.1f}%"
        return (p["id"], p["sector"][:4], f"{p['period']}mo", f"{p['start']:%Y-%m}",
                f"{p['end']:%Y-%m}", p["model"], f"{p['freq']}x", filter_label(p["X"], p["Y"]),
                p["status"], roi, spy, vs, shp, mdd, p["note"])

    def _safe_snapshot(self):
        """Snapshot the Run Filters, or None if a field is mid-edit / invalid."""
        try:
            return self._snapshot_filters()
        except Exception:
            return None

    def _perm_visible(self, p, snap=_UNSET) -> bool:
        """Whether permutation p should appear in the list under the current view filter
        (status checkboxes + optional 'Match Run Filters')."""
        var = self.view_status_vars.get(p["status"])
        if var is not None and not var.get():
            return False
        if self.view_match_filters.get():
            if snap is _UNSET:
                snap = self._safe_snapshot()
            if snap is not None and not perm_passes_filter(p, snap):
                return False
        return True

    def _fill_table(self):
        prev_sel = None
        sel = self.tree.selection()
        if sel:
            tags = self.tree.item(sel[0], "tags")
            if len(tags) >= 2:
                prev_sel = int(tags[1])
        self.tree.delete(*self.tree.get_children())
        self._iid.clear()
        snap = self._safe_snapshot() if self.view_match_filters.get() else _UNSET
        for p in self.perms:
            if not self._perm_visible(p, snap):
                continue
            iid = self.tree.insert("", "end", values=self._row_vals(p), tags=(p["status"], str(p["id"])))
            self._iid[p["id"]] = iid
        if prev_sel is not None and prev_sel in self._iid:
            self.tree.selection_set(self._iid[prev_sel])

    def _view_completed_only(self):
        """Quick action: show only completed (Done) permutations."""
        for st, v in self.view_status_vars.items():
            v.set(st == ST_DONE)
        self.view_match_filters.set(False)
        self._fill_table()
        done = sum(1 for p in self.perms if p["status"] == ST_DONE)
        self._act(f"[View] Showing {done} completed permutation(s) only.")

    def _reset_view(self):
        """Show every permutation regardless of status or Run Filters."""
        for v in self.view_status_vars.values():
            v.set(True)
        self.view_match_filters.set(False)
        self._fill_table()
        self._act("[View] Showing all permutations.")

    def _wire_view_traces(self):
        """Refresh the list when Run Filters change, but only while 'Match Run Filters' is on."""
        def cb(*_):
            if self.view_match_filters.get():
                try:
                    self._fill_table()
                except Exception:
                    pass
        for v in ([self.id_from, self.id_to, self.date_from, self.date_to]
                  + list(self.sector_vars.values()) + list(self.model_vars.values())
                  + list(self.period_vars.values()) + list(self.freq_vars.values())
                  + list(self.filt_vars.values())):
            v.trace_add("write", cb)

    def _update_row(self, pid):
        p = next((x for x in self.perms if x["id"] == pid), None)
        if not p:
            return
        # A status change can move a row into/out of the active view filter
        # (e.g. Running -> Done). Rebuild the list when its membership flips.
        if self._perm_visible(p) != (pid in self._iid):
            self._fill_table()
            return
        if pid in self._iid:
            self.tree.item(self._iid[pid], values=self._row_vals(p), tags=(p["status"], str(pid)))

    def _show_log(self, _evt=None):
        sel = self.tree.selection()
        if not sel:
            return
        tags = self.tree.item(sel[0], "tags")
        if len(tags) < 2:
            return
        p = next((x for x in self.perms if x["id"] == int(tags[1])), None)
        self.log_box.configure(state="normal")
        self.log_box.delete("1.0", "end")
        self.log_box.insert("end", "".join(p["log"]) if (p and p["log"]) else "(no log yet)")
        self.log_box.configure(state="disabled")
        self.log_box.see("end")

    def _open_perm_json(self, event=None):
        """Double-click a permutation row -> open its saved result JSON in the OS default app."""
        import glob
        import subprocess
        tree = event.widget if event is not None else self.tree
        row = tree.identify_row(event.y) if event is not None else (
            tree.selection()[0] if tree.selection() else "")
        if not row:
            return
        tags = tree.item(row, "tags")
        if len(tags) < 2:
            return
        pid = int(tags[1])
        p = next((x for x in self.perms if x["id"] == pid), None)
        if p is None:
            return
        path = p.get("result_path")
        if not path or not os.path.exists(path):
            # fallback: newest matching JSON anywhere under the results dir
            matches = glob.glob(os.path.join(RESULTS_DIR, "**", f"psperm_{pid:04d}_*.json"),
                                recursive=True)
            path = max(matches, key=os.path.getmtime) if matches else None
        if not path:
            self.messagebox.showinfo(
                "No result yet",
                f"Permutation #{pid} has no saved JSON yet.\n({p['status']})")
            return
        try:
            if sys.platform.startswith("win"):
                os.startfile(path)                       # type: ignore[attr-defined]
            elif sys.platform == "darwin":
                subprocess.Popen(["open", path])
            else:
                subprocess.Popen(["xdg-open", path])
            self._act(f"[Open] {path}")
        except Exception as e:
            self.messagebox.showerror("Open failed", f"Could not open:\n{path}\n\n{e}")

    # ---- worker pop-out windows ---------------------------------------- #
    def _worker_window(self, widx):
        tk, ttk = self.tk, self.ttk
        win = self._worker_windows.get(widx)
        if win and win["top"].winfo_exists():
            return win
        top = tk.Toplevel(self.root)
        top.title(f"Worker {widx} — live log")
        top.geometry(f"580x360+{70 + widx*36}+{70 + widx*36}")
        hdr = ttk.Label(top, text=f"Worker {widx}", font=("Segoe UI", 9, "bold"))
        hdr.pack(anchor="w", padx=4, pady=(4, 0))
        txt = self.scrolledtext.ScrolledText(top, wrap="word", font=("Consolas", 9))
        txt.pack(fill="both", expand=True)
        self._worker_windows[widx] = {"top": top, "text": txt, "hdr": hdr}
        return self._worker_windows[widx]

    # ---- stats / logs -------------------------------------------------- #
    def _refresh_stats(self):
        c = {ST_PENDING: 0, ST_SKIPPED: 0, ST_RUNNING: 0, ST_DONE: 0, ST_ERROR: 0}
        for p in self.perms:
            c[p["status"]] = c.get(p["status"], 0) + 1
        self.stats_var.set(f"P{c[ST_PENDING]} S{c[ST_SKIPPED]} R{c[ST_RUNNING]} "
                           f"D{c[ST_DONE]} E{c[ST_ERROR]} / {len(self.perms)}")

    def _act(self, text):
        self.act_box.configure(state="normal")
        self.act_box.insert("end", text + "\n")
        self.act_box.see("end")
        self.act_box.configure(state="disabled")

    def _load_existing(self):
        subdir = _safe_results_subdir(self.folder_var.get())
        files, runs = load_json_results_into_perms(self.perms, subdir)
        for p in self.perms:
            if p["results"]:
                if p["status"] == ST_ERROR:
                    continue                      # keep the overspend-error note from the loader
                r = p["result"]; roi = (r.get("total_roi_pct") if r else None)
                p["note"] = f"{len(p['results'])} run(s)" + (f"  ROI {roi:+.2f}%" if roi is not None else "")
        where = f"results/{subdir}" if subdir else "results root"
        if runs:
            self._act(f"[Startup] Restored {runs} run(s) from {files} file(s) "
                      f"in {where}.")
        else:
            self._act(f"[Startup] No prior runs found in {where} — "
                      f"every selected cell will run from scratch.")


    def _apply_preset_filters(self, preset: dict):
        """Tick exactly the cells the preset names; leave anything unspecified alone."""
        f = preset.get("filters") or {}
        name = preset.get("name", "preset")

        def _set(varmap, wanted, label, key=lambda x: x):
            if wanted is None:
                return
            want = {key(w) for w in wanted}
            hit = 0
            for k, var in varmap.items():
                on = key(k) in want if not isinstance(k, tuple) else k in want
                var.set(on)
                hit += int(on)
            if hit == 0:
                self._act(f"[preset] WARNING: no {label} matched {sorted(want)}")
            else:
                self._act(f"[preset] {label}: {hit} selected")

        _set(self.sector_vars, f.get("sectors"), "sectors")
        _set(self.model_vars, f.get("models"), "models")
        _set(self.period_vars, f.get("periods"), "periods", key=int)
        _set(self.freq_vars, f.get("freqs"), "cadences", key=int)
        if f.get("filters") is not None:
            want = {parse_filter_loose(x) for x in f["filters"]}
            for k, var in self.filt_vars.items():
                var.set(k in want)
            self._act(f"[preset] filter configs: {sorted(want)}")
        if f.get("date_from"):
            self.date_from.set(f["date_from"])
        if f.get("date_to"):
            self.date_to.set(f["date_to"])
        if f.get("id_from"):
            self.id_from.set(int(f["id_from"]))
        if f.get("id_to"):
            self.id_to.set(int(f["id_to"]))

        self._act(f"[preset] loaded '{name}' — "
                  f"feedback={'ON' if self.fb_var.get() else 'OFF'}, "
                  f"screened_news={'ON' if self.screen_var.get() else 'OFF'}, "
                  f"orders={self.ordmode_var.get()}, "
                  f"min_articles={self._min_articles()}, "
                  f"folder={self.folder_var.get() or '(root)'}, "
                  f"runs/perm={self.runs_var.get()}")
        if self.dry_var.get():
            self._act("[preset] DRY-RUN is on — untick it to spend API budget.")
        self._refresh_stats()


    def _apply_preset_options(self, preset: dict):
        """Push a preset's `options` straight onto the live tk variables.

        __init__ folds options into self.cfg before the vars are constructed; this
        does the same job after they exist, so loading a preset from the button
        behaves identically to launching with --preset.
        """
        o = preset.get("options") or {}
        targets = {
            "feedback": self.fb_var, "screened_news": self.screen_var,
            "min_articles": self.minart_var, "min_articles_days": self.minartdays_var,
            "txn_cost_bps": self.txn_var, "order_mode": self.ordmode_var,
            "rebalance_mode": self.rebmode_var,
            "prompt_framing": self.framing_var,
            "screened_corpus": self.corpus_var,
            "results_subdir": self.folder_var, "runs_per_perm": self.runs_var,
            "threads": self.threads_var, "budget": self.budget_var,
            "risk_free": self.rf_var, "whole_shares": self.whole_var,
            "dry_run": self.dry_var,
        }
        for key, var in targets.items():
            if key in o:
                try:
                    var.set(o[key])
                except Exception as e:              # noqa: BLE001
                    self._act(f"[preset] could not set {key}={o[key]!r}: {e}")
            elif key in EXPERIMENT_DEFINING_DEFAULTS:
                # Not specified by the preset -> force the default instead of
                # inheriting a stale value from the config file or a prior run.
                dflt = EXPERIMENT_DEFINING_DEFAULTS[key]
                try:
                    if var.get() != dflt:
                        self._act(f"[preset] {key} unspecified — forcing default "
                                  f"{dflt!r} (was {var.get()!r})")
                    var.set(dflt)
                except Exception as e:              # noqa: BLE001
                    self._act(f"[preset] could not default {key}: {e}")

    def _on_load_preset(self):
        """Pick a preset JSON and apply it to the current session."""
        if self._worker is not None and self._worker.is_alive():
            self.messagebox.showwarning(
                "Run in progress",
                "Stop the current run before loading a preset.")
            return
        initial = os.path.join(ROOT, "presets")
        path = self.filedialog.askopenfilename(
            title="Load preset",
            initialdir=initial if os.path.isdir(initial) else ROOT,
            filetypes=[("Preset JSON", "*.json"), ("All files", "*.*")])
        if not path:
            return
        try:
            preset = load_preset(path)
        except Exception as e:                       # noqa: BLE001
            self.messagebox.showerror("Bad preset", f"Could not load:\n{path}\n\n{e}")
            self._act(f"[preset] FAILED to load {os.path.basename(path)}: {e}")
            return
        self.preset = preset
        self._apply_preset_options(preset)
        self._apply_preset_filters(preset)
        self._persist()
        self.root.title("Portfolio-Sim Permutation Runner \u2014 "
                        f"{preset.get('name', os.path.basename(path))}")


    def _begin_autostart(self, remaining: "Optional[int]" = None):
        """Count down, then press Start. Announces LIVE vs DRY-RUN first.

        The countdown is the abort window: Stop, or closing the window, cancels it.
        autostart is deliberately CLI-only and never a preset field, so a preset file
        can never cause unattended API spending on its own.
        """
        if remaining is None:
            remaining = self.autostart_delay
            mode = ("DRY-RUN (mock backend, no API calls)" if self.dry_var.get()
                    else "LIVE — this will spend API budget")
            self._act(f"[autostart] {mode}")
            self._act(f"[autostart] starting in {remaining}s — close this window to abort")
        if self._worker is not None and self._worker.is_alive():
            self._act("[autostart] a run is already active; not starting again")
            return
        if remaining <= 0:
            self._act("[autostart] starting now")
            self._on_start()
            return
        if remaining <= 3 or remaining % 5 == 0:
            self._act(f"[autostart] {remaining}...")
        self.root.after(1000, lambda: self._begin_autostart(remaining - 1))

    # ---- controls ------------------------------------------------------ #
    def _set_filters_enabled(self, enabled):
        state = "normal" if enabled else "disabled"
        for w in self._filter_widgets:
            try:
                w.configure(state=state if not isinstance(w, self.ttk.Combobox) else ("readonly" if enabled else "disabled"))
            except Exception:
                pass

    def _min_articles_days(self) -> int:
        """Recency window for the density gate; blank/garbage means 90."""
        try:
            return max(1, int(self.minartdays_var.get()))
        except Exception:
            return 90

    def _min_articles(self) -> int:
        """Min-articles as a non-negative int; blank/garbage means 0 (no gate)."""
        try:
            return max(0, int(self.minart_var.get()))
        except Exception:
            return 0

    def _txn_bps(self) -> float:
        """Txn bps as a non-negative float; a blank/garbage field means 0 (frictionless)."""
        try:
            return max(0.0, float(self.txn_var.get()))
        except Exception:
            return 0.0

    def _persist(self):
        safe_folder = _safe_results_subdir(self.folder_var.get())
        self.folder_var.set(safe_folder)          # reflect sanitized value in the field
        self.cfg.update({
            "openai_api_key": self.key_var.get().strip(),
            "nyt_api_key": self.nyt_var.get().strip(),
            "results_subdir": safe_folder,
            "threads": self.threads_var.get(), "runs_per_perm": self.runs_var.get(),
            "budget": self.budget_var.get(), "risk_free": self.rf_var.get(),
            "whole_shares": self.whole_var.get(), "dry_run": self.dry_var.get(),
            "feedback": self.fb_var.get(), "txn_cost_bps": self._txn_bps(),
            "screened_news": self.screen_var.get(),
            "min_articles": self._min_articles(),
            "min_articles_days": self._min_articles_days(),
            "order_mode": self.ordmode_var.get(),
            "rebalance_mode": self.rebmode_var.get(),
            "prompt_framing": self.framing_var.get(),
            "screened_corpus": self.corpus_var.get().strip(),
        })
        save_app_config(self.cfg)

    def _scheduled_map(self, filters, target_runs, retry_errors_only):
        """{perm_id: runs_scheduled} for a Start/Retry right now (mirrors RunnerThread)."""
        m = {}
        for p in self.perms:
            if retry_errors_only:
                if p["status"] == ST_ERROR and perm_passes_filter(p, filters):
                    m[p["id"]] = 1
            else:
                if p["status"] == ST_SKIPPED or not perm_passes_filter(p, filters):
                    continue
                n = max(0, target_runs - len(p.get("results", [])))
                if n > 0:
                    m[p["id"]] = n
        return m

    def _count_scheduled(self, filters, target_runs, retry_errors_only):
        return sum(self._scheduled_map(filters, target_runs, retry_errors_only).values())

    # ── "Queued runs" window: only the permutations in the CURRENT batch ──────
    def _open_batch_window(self, sched_map):
        tk, ttk = self.tk, self.ttk
        self._batch_sched = dict(sched_map)
        self._batch_done = {}
        bw = getattr(self, "_batch_win", None)
        if bw is None or not bw["top"].winfo_exists():
            top = tk.Toplevel(self.root)
            top.title("Queued runs — current batch")
            top.geometry("760x480+120+90")
            hdr = ttk.Label(top, text="", font=("Segoe UI", 10, "bold"))
            hdr.pack(anchor="w", padx=6, pady=(6, 2))
            ttk.Label(top, text="Only the runs queued by the current Start/Retry.",
                      foreground="#666").pack(anchor="w", padx=6)
            cols = ("#", "Permutation", "Runs", "Status", "ROI", "S&P 500", "Sharpe")
            tree = ttk.Treeview(top, columns=cols, show="headings")
            widths = {"#": 50, "Permutation": 330, "Runs": 60, "Status": 74, "ROI": 84,
                      "S&P 500": 74, "Sharpe": 60}
            for c in cols:
                tree.heading(c, text=c)
                tree.column(c, width=widths[c], anchor="w" if c == "Permutation" else "center")
            ys = ttk.Scrollbar(top, orient="vertical", command=tree.yview)
            tree.configure(yscrollcommand=ys.set)
            ys.pack(side="right", fill="y")
            tree.pack(fill="both", expand=True)
            for st, bg in ROW_BG.items():
                if bg:
                    tree.tag_configure(st, background=bg)
            tree.bind("<Double-1>", self._open_perm_json)
            self._batch_win = {"top": top, "tree": tree, "hdr": hdr, "iid": {}}
        bw = self._batch_win
        bw["tree"].delete(*bw["tree"].get_children())
        bw["iid"].clear()
        for p in self.perms:
            if p["id"] in self._batch_sched:
                iid = bw["tree"].insert("", "end", values=self._batch_row(p),
                                        tags=(p["status"], str(p["id"])))
                bw["iid"][p["id"]] = iid
        self._update_batch_header()
        bw["top"].deiconify(); bw["top"].lift()

    def _batch_row(self, p):
        r = p.get("result") or {}
        rmx = (r.get("risk_metrics") or {}) if r else {}
        bench = (r.get("benchmark_risk_metrics") or {}) if r else {}
        roi = "" if r.get("total_roi_pct") is None else f"{r['total_roi_pct']:+.2f}%"
        spy = "" if bench.get("total_return") is None else f"{bench['total_return']*100:+.2f}%"
        shp = "" if rmx.get("sharpe") is None else f"{rmx['sharpe']:.2f}"
        done = self._batch_done.get(p["id"], 0)
        sched = self._batch_sched.get(p["id"], 0)
        return (p["id"], perm_label(p), f"{done}/{sched}", p["status"], roi, spy, shp)

    def _update_batch_row(self, pid):
        bw = getattr(self, "_batch_win", None)
        if not bw or not bw["top"].winfo_exists() or pid not in bw["iid"]:
            return
        p = next((x for x in self.perms if x["id"] == pid), None)
        if p:
            bw["tree"].item(bw["iid"][pid], values=self._batch_row(p),
                            tags=(p["status"], str(pid)))

    def _update_batch_header(self):
        bw = getattr(self, "_batch_win", None)
        if bw and bw["top"].winfo_exists():
            remaining = self._run_total - self._run_finished
            bw["hdr"].config(
                text=f"Runs completed: {self._run_finished} / {self._run_total}   "
                     f"({remaining} remaining · {len(self._batch_sched)} permutations)")

    def _reopen_batch_window(self):
        if getattr(self, "_batch_sched", None):
            self._open_batch_window(self._batch_sched)
        else:
            self.messagebox.showinfo("No batch yet", "Start a run first to see its queued runs.")

    def _start_worker(self, retry_errors_only):
        if self._worker and self._worker.is_alive():
            return
        if not self.dry_var.get() and not self.key_var.get().strip():
            self.messagebox.showwarning("No API key", "Enter an OpenAI API key or enable Dry-run.")
            return
        self._persist()
        self._stop.clear(); self._pause.clear()
        filters = self._snapshot_filters()
        sched_map = self._scheduled_map(filters, self.runs_var.get(), retry_errors_only)
        n_queued = sum(sched_map.values())
        if n_queued == 0:
            self.messagebox.showinfo("Nothing to run",
                "No permutations match the current filters (all skipped/complete, or none selected).")
            return
        self._run_total = n_queued; self._run_finished = 0
        self.run_prog_var.set(f"Runs: {n_queued} queued / 0 completed")
        self.pb_var.set(0.0)
        self._act(f"[Start] {n_queued} run(s) queued across {len(sched_map)} permutation(s).  "
                  f"feedback={'ON' if self.fb_var.get() else 'OFF'}  txn={self._txn_bps():g}bps  "
                  f"screened_news={'ON' if self.screen_var.get() else 'OFF'}  "
                  f"min_articles={self._min_articles()}/{self._min_articles_days()}d  "
                  f"orders={self.ordmode_var.get()}  "
                  f"folder={self.folder_var.get() or '(root)'}")
        self._open_batch_window(sched_map)
        self._worker = RunnerThread(
            self.perms, self._q, api_key=self.key_var.get().strip(),
            dry_run=self.dry_var.get(), budget=self.budget_var.get(),
            whole_shares=self.whole_var.get(), rf=self.rf_var.get(),
            stop_evt=self._stop, pause_evt=self._pause, n_threads=self.threads_var.get(),
            filters=filters, target_runs=self.runs_var.get(),
            retry_errors_only=retry_errors_only, results_subdir=self.folder_var.get(),
            feedback=self.fb_var.get(), txn_cost_bps=self._txn_bps(),
            screened_news=self.screen_var.get(), min_articles=self._min_articles(),
            order_mode=self.ordmode_var.get(),
            rebalance_mode=self.rebmode_var.get(),
            prompt_framing=self.framing_var.get(),
            screened_corpus=self.corpus_var.get().strip(),
            min_articles_days=self._min_articles_days())
        self.b_start.config(state="disabled"); self.b_retry.config(state="disabled")
        self.b_retry_all.config(state="disabled")
        self.b_pause.config(state="normal"); self.b_stop.config(state="normal")
        self._set_filters_enabled(False)
        self.folder_entry.config(state="disabled")
        self.fb_check.config(state="disabled"); self.txn_entry.config(state="disabled")
        self.screen_check.config(state="disabled"); self.minart_entry.config(state="disabled")
        self.ordmode_combo.config(state="disabled")
        self.minartdays_entry.config(state="disabled")
        self.b_preset.config(state="disabled")
        self._worker.start()

    def _on_start(self):
        self._start_worker(retry_errors_only=False)

    def _on_retry(self):
        self._start_worker(retry_errors_only=True)

    def _on_retry_all(self):
        """Re-run EVERY filter-matched permutation regardless of status, replacing prior
        results (clears in-memory results and deletes their old JSON files first)."""
        if self._worker and self._worker.is_alive():
            return
        import glob
        filters = self._snapshot_filters()
        matched = [p for p in self.perms
                   if p["status"] != ST_SKIPPED and perm_passes_filter(p, filters)]
        if not matched:
            self.messagebox.showinfo("Nothing to run",
                                     "No permutations match the current filters.")
            return
        files = []
        for p in matched:
            files += glob.glob(os.path.join(RESULTS_DIR, "**", f"psperm_{p['id']:04d}_*.json"),
                               recursive=True)
        n_runs = self.runs_var.get() * len(matched)
        if not self.messagebox.askyesno(
                "Retry all (replace)",
                f"Re-run ALL {len(matched)} matching permutation(s) as {n_runs} fresh run(s), "
                f"regardless of status, REPLACING their previous results.\n\n"
                f"This permanently deletes {len(files)} existing result file(s) for those "
                f"permutations. Continue?"):
            return
        removed = 0
        for fp in files:
            try:
                os.remove(fp); removed += 1
            except OSError:
                pass
        for p in matched:
            p["results"] = []
            p["result"] = None
            p["result_path"] = None
            p["status"] = ST_PENDING
            p["note"] = ""
            self._update_row(p["id"])
        self._act(f"[Retry All] Deleted {removed} old file(s); re-running "
                  f"{len(matched)} permutation(s) as {n_runs} run(s).")
        self._refresh_stats()
        self._start_worker(retry_errors_only=False)

    def _on_pause(self):
        if self._pause.is_set():
            self._pause.clear(); self.b_pause.config(text="⏸ Pause")
        else:
            self._pause.set(); self.b_pause.config(text="▶ Resume")

    def _on_stop(self):
        self._stop.set(); self._pause.clear()
        self.b_stop.config(state="disabled")
        self.b_pause.config(state="disabled")
        self.cur_var.set("Stopping…  (no new runs; in-flight runs abort at next step)")
        self._act("[Stop] Stopping — queue halted; in-flight runs will abort shortly.")

    def _on_excel(self):
        err = save_excel(self.perms, EXCEL_PATH)
        self._act(f"[Excel] {'ERROR: ' + err if err else 'saved ' + EXCEL_PATH}")

    # ---- queue drain (main thread) ------------------------------------- #
    def _drain(self):
        try:
            while True:
                msg = self._q.get_nowait()
                k = msg["kind"]
                if k == "plan":
                    self._run_total = msg["total"]; self._run_finished = 0
                    self.pb_var.set(0.0)
                    self.run_prog_var.set(f"Runs: {msg['total']} queued / 0 completed")
                    self._act(f"[Plan] {msg['total']} run(s) scheduled.")
                elif k == "status":
                    p = next((x for x in self.perms if x["id"] == msg["id"]), None)
                    if p is not None:
                        p["status"] = msg["status"]; p["note"] = msg.get("note", p["note"])
                        if msg.get("result") is not None:
                            p["result"] = msg["result"]; p["results"].append(msg["result"])
                        if msg.get("worker") is not None:
                            p["worker"] = msg["worker"]
                        if msg.get("json_path"):
                            p["result_path"] = msg["json_path"]
                        self._update_row(msg["id"])
                        if msg["status"] == ST_RUNNING:
                            self.cur_var.set(perm_label(p))
                            win = self._worker_window(msg.get("worker", 0))
                            win["hdr"].config(text=f"Worker {msg.get('worker',0)}  ▶  {perm_label(p)}")
                            win["text"].insert("end", f"\n{'='*60}\n{perm_label(p)}\n{'='*60}\n"); win["text"].see("end")
                        elif msg["status"] in (ST_DONE, ST_ERROR):
                            self._run_finished += 1
                            if msg["id"] in self._batch_sched:
                                self._batch_done[msg["id"]] = self._batch_done.get(msg["id"], 0) + 1
                            if self._run_total:
                                self.pb_var.set(100.0 * self._run_finished / self._run_total)
                            self.run_prog_var.set(
                                f"Runs: {self._run_total} queued / {self._run_finished} completed")
                        self._update_batch_row(msg["id"])
                        self._update_batch_header()
                        self._refresh_stats()
                elif k == "perm_log":
                    p = next((x for x in self.perms if x["id"] == msg["id"]), None)
                    win = self._worker_windows.get(msg.get("worker"))
                    if win and win["top"].winfo_exists():
                        win["text"].insert("end", msg["text"]); win["text"].see("end")
                    sel = self.tree.selection()
                    if sel:
                        tags = self.tree.item(sel[0], "tags")
                        if len(tags) >= 2 and int(tags[1]) == msg["id"]:
                            self.log_box.configure(state="normal")
                            self.log_box.insert("end", msg["text"]); self.log_box.see("end")
                            self.log_box.configure(state="disabled")
                elif k == "log":
                    self._act(msg["text"])
                elif k == "done":
                    self.b_start.config(state="normal"); self.b_retry.config(state="normal")
                    self.b_retry_all.config(state="normal")
                    self.b_pause.config(state="disabled", text="⏸ Pause")
                    self.b_stop.config(state="disabled")
                    self._set_filters_enabled(True)
                    self.folder_entry.config(state="normal")
                    self.fb_check.config(state="normal"); self.txn_entry.config(state="normal")
                    self.screen_check.config(state="normal")
                    self.minart_entry.config(state="normal")
                    self.ordmode_combo.config(state="readonly")
                    self.minartdays_entry.config(state="normal")
                    self.b_preset.config(state="normal")
                    was_stopped = self._stop.is_set()
                    self.cur_var.set("Stopped" if was_stopped else "Finished")
                    tail = "(stopped)" if was_stopped else "(done)"
                    self.run_prog_var.set(
                        f"Runs: {self._run_total} queued / {self._run_finished} completed  {tail}")
                    self._act("[Stop] Batch stopped." if was_stopped else "[Done] Batch complete.")
                    err = save_excel(self.perms, EXCEL_PATH)
                    self._act(f"[Excel] {'ERROR: ' + err if err else 'saved ' + EXCEL_PATH}")
        except queue.Empty:
            pass
        self.root.after(200, self._drain)

    # ---- sort ---------------------------------------------------------- #
    def _sort(self, col):
        rev = self._sort_rev.get(col, False)

        def key(item):
            v = self.tree.set(item, col)
            m = re.search(r"-?\d+\.?\d*", v)
            return float(m.group()) if m else v.lower()
        items = list(self.tree.get_children(""))
        try:
            items.sort(key=key, reverse=rev)
        except Exception:
            items.sort(key=lambda i: self.tree.set(i, col), reverse=rev)
        for pos, it in enumerate(items):
            self.tree.move(it, "", pos)
        self._sort_rev[col] = not rev

    def run(self):
        self.root.mainloop()


# =========================================================================== #
#  Headless + CLI                                                             #
# =========================================================================== #
def _apply_cli_filters(perms, args, preset=None):
    pf = (preset or {}).get("filters") or {}
    if pf:
        from dateutil.parser import parse as _dp
        # A bare "YYYY-MM" must mean the FIRST of that month. dateutil defaults the
        # missing day to today's day-of-month, which silently drops the window whose
        # start is the 1st.
        def _p(x):
            return _dp(str(x), default=datetime(2000, 1, 1))
        return {
            "id_from": int(pf.get("id_from", 1)),
            "id_to": int(pf.get("id_to", max(p["id"] for p in perms))),
            "sectors": set(pf.get("sectors") or SECTORS),
            "models": set(pf.get("models") or MODELS),
            "periods": {int(x) for x in (pf.get("periods") or PERIOD_MONTHS)},
            "freqs": {int(x) for x in (pf.get("freqs") or FREQS)},
            "filters": {parse_filter_loose(x)
                        for x in (pf.get("filters") or FILTER_CONFIGS)},
            "date_from": _p(pf["date_from"]) if pf.get("date_from") else START_DATES[0],
            "date_to": _p(pf["date_to"]) if pf.get("date_to") else START_DATES[-1],
        }
    return {
        "id_from": 1, "id_to": max(p["id"] for p in perms),
        "sectors": set(args.sectors.split(",")) if args.sectors else set(SECTORS),
        "models": set(args.models.split(",")) if args.models else set(MODELS),
        "periods": {int(x) for x in args.periods.split(",")} if args.periods else set(PERIOD_MONTHS),
        "freqs": {int(x) for x in args.freqs.split(",")} if args.freqs else set(FREQS),
        "filters": {parse_filter(x) for x in args.filters.split(",")} if args.filters else set(FILTER_CONFIGS),
        "date_from": START_DATES[0], "date_to": START_DATES[-1],
    }


def run_headless(args, preset=None):
    perms = build_permutations()
    o = (preset or {}).get("options") or {}
    for attr, key in (("runs", "runs_per_perm"), ("budget", "budget"),
                      ("workers", "threads"), ("folder", "results_subdir"),
                      ("min_articles", "min_articles"),
                      ("min_articles_days", "min_articles_days")):
        if key in o:
            setattr(args, attr, o[key])
    if "risk_free" in o:
        args.risk_free = o["risk_free"]
    if "feedback" in o:
        args.feedback = "on" if o["feedback"] else "off"
    if "screened_news" in o:
        args.screened_news = "on" if o["screened_news"] else "off"
    if "order_mode" in o:
        args.order_mode = o["order_mode"]
    if "rebalance_mode" in o:
        args.rebalance_mode = o["rebalance_mode"]
    if "prompt_framing" in o:
        args.prompt_framing = o["prompt_framing"]
    if "screened_corpus" in o:
        args.screened_corpus = o["screened_corpus"]
    if "txn_cost_bps" in o:
        args.txn_bps = o["txn_cost_bps"]
    load_json_results_into_perms(perms, _safe_results_subdir(args.folder))
    q: "queue.Queue[dict]" = queue.Queue()
    stop, pause = threading.Event(), threading.Event()
    rt = RunnerThread(perms, q, api_key=os.environ.get("OPENAI_API_KEY", ""),
                      dry_run=not args.live, budget=args.budget, whole_shares=args.whole_shares,
                      rf=args.risk_free, stop_evt=stop, pause_evt=pause, n_threads=args.workers,
                      filters=_apply_cli_filters(perms, args, preset), target_runs=args.runs,
                      results_subdir=args.folder,
                      feedback=(args.feedback != "off"), txn_cost_bps=args.txn_bps,
                      screened_news=(args.screened_news == "on"),
                      min_articles=args.min_articles, order_mode=args.order_mode,
                      rebalance_mode=args.rebalance_mode,
                      prompt_framing=args.prompt_framing,
                      screened_corpus=args.screened_corpus,
                      min_articles_days=args.min_articles_days)
    print("Running permutations headless...")
    by_id = {p["id"]: p for p in perms}
    rt.start()
    done = 0
    while True:
        msg = q.get()
        if msg["kind"] == "status" and msg["status"] in (ST_DONE, ST_ERROR):
            done += 1
            p = by_id.get(msg["id"])
            if p is not None:
                p["status"] = msg["status"]; p["note"] = msg.get("note", p["note"])
                if msg.get("result") is not None:
                    p["result"] = msg["result"]; p["results"].append(msg["result"])
            print(f"[{done}] #{msg['id']:04d} {msg['status']}: {msg.get('note','')}")
        elif msg["kind"] == "done":
            break
    err = save_excel(perms, EXCEL_PATH)
    print("Excel:", "ERROR " + err if err else EXCEL_PATH)
    if args.out:
        with open(args.out, "w", encoding="utf-8") as f:
            json.dump([{"id": p["id"], "label": perm_label(p), "status": p["status"],
                        "result": p["result"]} for p in perms
                       if p["results"] or p["status"] == ST_ERROR],
                      f, indent=2, default=str)
        print("Wrote", args.out)


def main(argv=None):
    ap = argparse.ArgumentParser(description="Portfolio-sim permutation runner + GUI")
    ap.add_argument("--preset", default="",
                    help="JSON preset that preloads run filters + option knobs "
                         "(see presets/). Opens the GUI already configured.")
    ap.add_argument("--autostart", action="store_true",
                    help="open the GUI and press Start automatically. The window stays "
                         "up with full progress/pause/stop; you just do not have to click.")
    ap.add_argument("--autostart-delay", type=int, default=8, dest="autostart_delay",
                    help="seconds to wait before auto-starting, so you can abort (default 8)")
    ap.add_argument("--headless", action="store_true")
    ap.add_argument("--save-excel-from-json", action="store_true")
    ap.add_argument("--live", action="store_true", help="use RealBackend (OpenAI); default dry-run")
    ap.add_argument("--workers", type=int, default=1)
    ap.add_argument("--runs", type=int, default=1)
    ap.add_argument("--budget", type=float, default=BUDGET_DEFAULT)
    ap.add_argument("--whole-shares", action="store_true")
    ap.add_argument("--risk-free", type=float, default=0.0)
    ap.add_argument("--folder", default="")
    ap.add_argument("--feedback", choices=["on", "off"], default="on",
                    help="self-reflection loop; 'off' is the control arm")
    ap.add_argument("--txn-bps", type=float, default=0.0, dest="txn_bps",
                    help="transaction cost in bps on round-trip notional per rebalance")
    ap.add_argument("--screened-news", choices=["on", "off"], default="off",
                    dest="screened_news",
                    help="Stage One reads tech_screened_corpus.csv (Paper B only)")
    ap.add_argument("--order-mode", choices=["dollars", "weights"], default="dollars",
                    dest="order_mode",
                    help="weights: model emits target fractions; the engine does the "
                         "arithmetic (overdraw impossible, no re-prompt loop)")
    ap.add_argument("--rebalance-mode",
                    choices=["reset", "hold", "hold_redeploy"], default="reset",
                    dest="rebalance_mode",
                    help="reset: liquidate the book each rebalance and re-buy to target "
                         "weights (original). hold: trade only the delta and leave kept "
                         "positions alone so winners keep compounding (needs weights mode)")
    ap.add_argument("--prompt-framing",
                    choices=["reset", "incumbent", "incumbent_invested"], default="reset",
                    dest="prompt_framing",
                    help="reset: the rebalance prompt says the book was already sold and "
                         "asks the model to rebuild from scratch (original). incumbent: "
                         "the prompt states the book still held, the gap to the next "
                         "decision, and a switching hurdle. incumbent_invested: incumbent "
                         "plus an explicit stay-invested instruction (incumbent alone drove "
                         "mean cash from 1.5%% to 7.1%% because the hurdle suppressed buying "
                         "as well as selling)")
    ap.add_argument("--screened-corpus", default="", dest="screened_corpus",
                    help="alternate screened corpus for Stage One, e.g. "
                         "tech_scored_ge7.csv from score_news_relevance.py --emit. "
                         "Empty uses tech_screened_corpus.csv. Needs --screened-news on")
    ap.add_argument("--min-articles-days", type=int, default=90, dest="min_articles_days",
                    help="recency window for the density gate (default 90)")
    ap.add_argument("--min-articles", type=int, default=0, dest="min_articles",
                    help="skip cells whose thinnest decision sees fewer than N articles "
                         "in the prior 90 days (0 = no gate)")
    ap.add_argument("--sectors"); ap.add_argument("--models"); ap.add_argument("--periods")
    ap.add_argument("--freqs"); ap.add_argument("--filters"); ap.add_argument("--out")
    args = ap.parse_args(argv)

    if args.save_excel_from_json:
        perms = build_permutations()
        _, runs = load_json_results_into_perms(perms)
        err = save_excel(perms, EXCEL_PATH)
        print(f"Loaded {runs} run(s). Excel:", "ERROR " + err if err else EXCEL_PATH)
        return 0
    preset = None
    if args.preset:
        try:
            preset = load_preset(args.preset)
        except Exception as e:
            print(f"Could not load preset {args.preset}: {e}")
            return 1
        print(f"Loaded preset: {preset.get('name', args.preset)}")
        if preset.get("description"):
            print(f"  {preset['description']}")
    if args.headless:
        run_headless(args, preset)
        return 0
    App(preset, autostart=args.autostart,
        autostart_delay=args.autostart_delay).run()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())